from dotenv import load_dotenv
from pathlib import Path

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

import os
import io
import uuid
import logging
import asyncpg
import bcrypt
import jwt
import random
import openpyxl
from openpyxl.styles import Font
from decimal import Decimal
from datetime import datetime, timezone, timedelta, date
from typing import List, Optional, Literal

from fastapi import FastAPI, APIRouter, HTTPException, Request, Response, Depends, Query
from fastapi import UploadFile, File, Form
from fastapi.responses import StreamingResponse
from starlette.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field, EmailStr

import db as _db

# ---------- App ----------
app = FastAPI(title="FrostCore API")
api = APIRouter(prefix="/api")

JWT_ALGO = "HS256"
SHUTTLE_ZONE = "SHUTTLE_ZONE_A"
SCHEMA_VERSION = 5  # bump to wipe + re-seed


# ---------- Helpers ----------
def jwt_secret() -> str:
    return os.environ["JWT_SECRET"]


def hash_pw(p: str) -> str:
    return bcrypt.hashpw(p.encode(), bcrypt.gensalt()).decode()


def verify_pw(plain: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(plain.encode(), hashed.encode())
    except Exception:
        return False


def create_token(user_id: str, email: str, ttl_minutes: int = 60 * 24 * 7) -> str:
    payload = {
        "sub": user_id,
        "email": email,
        "exp": datetime.now(timezone.utc) + timedelta(minutes=ttl_minutes),
        "type": "access",
    }
    return jwt.encode(payload, jwt_secret(), algorithm=JWT_ALGO)


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _get_anthropic_client():
    """Return (client, model) using whichever API key is available.

    Priority:
      1. AI_INTEGRATIONS_ANTHROPIC_API_KEY  — Replit-managed proxy key.
         Uses AI_INTEGRATIONS_ANTHROPIC_BASE_URL and a proxy-compatible model alias.
      2. ANTHROPIC_API_KEY — user-supplied direct key.
         Always talks to api.anthropic.com regardless of any integration base-url
         setting, because the Replit proxy only accepts its own integration keys.
    """
    import anthropic as _anthropic
    integ_key = os.environ.get("AI_INTEGRATIONS_ANTHROPIC_API_KEY", "")
    direct_key = os.environ.get("ANTHROPIC_API_KEY", "")
    if integ_key:
        base_url = os.environ.get("AI_INTEGRATIONS_ANTHROPIC_BASE_URL", "https://api.anthropic.com")
        return (
            _anthropic.Anthropic(api_key=integ_key, base_url=base_url),
            "claude-sonnet-4-5",
        )
    if direct_key:
        # claude-sonnet-4-5 is the current model available on this account.
        return (
            _anthropic.Anthropic(api_key=direct_key),
            "claude-sonnet-4-5",
        )
    raise RuntimeError("No Anthropic API key configured. Set ANTHROPIC_API_KEY in Secrets.")


def _shuttle_code(lane: int, level: int, depth: int) -> str:
    return f"SZA-{lane:02d}-L{level:02d}-D{depth:02d}"


# ---------- Models ----------
class RegisterIn(BaseModel):
    email: EmailStr
    password: str
    name: str
    role: Literal["admin", "manager", "operator"] = "operator"


class LoginIn(BaseModel):
    email: EmailStr
    password: str


class SkuIn(BaseModel):
    sku_code: str
    name: str
    category: str
    bag_color: Optional[str] = None
    weight_per_bag: Optional[float] = None
    bags_per_pallet: Optional[int] = None
    dimensions: Optional[str] = None
    unit: str = "EA"
    unit_price: float = 0.0
    reorder_level: int = 10


BAG_COLORS = {"Green", "White", "Yellow"}

# Sensible packaging defaults by category, used for seeding new SKUs and for
# backfilling existing SKUs whose packaging columns were added after creation.
# Tuple: (bag_color, weight_per_bag_kg, bags_per_pallet, dimensions)
PACKAGING_DEFAULTS = {
    "Frozen Meat": ("White", 20.0, 48, "60x40x25 cm"),
    "Seafood":     ("White", 18.0, 50, "60x40x22 cm"),
    "Dairy":       ("Yellow", 25.0, 40, "60x40x30 cm"),
    "Ice Cream":   ("Yellow", 20.0, 45, "60x40x28 cm"),
    "Vegetables":  ("Green", 15.0, 60, "55x35x20 cm"),
    "Fruits":      ("Green", 12.0, 64, "55x35x18 cm"),
    "Bakery":      ("White", 10.0, 72, "50x30x20 cm"),
    "Ready Meals": ("White", 12.0, 60, "55x35x22 cm"),
    "Processed":   ("Yellow", 15.0, 56, "60x40x24 cm"),
}
DEFAULT_PACKAGING = ("White", 20.0, 48, "60x40x25 cm")


def _packaging_for(category: Optional[str]):
    """Return sensible (bag_color, weight_per_bag, bags_per_pallet, dimensions)
    for a category, falling back to a generic default for unknown categories."""
    return PACKAGING_DEFAULTS.get(category, DEFAULT_PACKAGING)


def _validate_sku_fields(body: "SkuIn"):
    # Bag Color is optional; normalise blank/whitespace to None so it stores as
    # NULL rather than an empty string.
    body.bag_color = body.bag_color.strip() if body.bag_color and body.bag_color.strip() else None
    if body.bag_color is not None and body.bag_color not in BAG_COLORS:
        raise HTTPException(400, "Bag Color must be one of: Green, White, Yellow")
    if body.weight_per_bag is None or body.weight_per_bag <= 0:
        raise HTTPException(400, "Weight per Bag must be greater than 0")
    if body.bags_per_pallet is None or body.bags_per_pallet <= 0:
        raise HTTPException(400, "Bags per Pallet must be greater than 0")



def _to_decimal(value):
    return Decimal(str(value)) if value is not None else None


# ---------- Excel import / export of SKUs ----------
# Editable SKU columns, in the order they appear in the spreadsheet. The export
# additionally appends a read-only "On Hand" column (ignored on import).
SKU_EXCEL_COLUMNS = [
    ("sku_code", "SKU Code"),
    ("name", "Name"),
    ("category", "Category"),
    ("bag_color", "Bag Color"),
    ("weight_per_bag", "Weight per Bag"),
    ("bags_per_pallet", "Bags per Pallet"),
    ("dimensions", "Dimensions"),
    ("unit", "Unit"),
    ("unit_price", "Unit Price"),
    ("reorder_level", "Reorder Level"),
]


def _excel_str(v):
    if v is None:
        return None
    v = str(v).strip()
    return v or None


def _excel_float(v, field):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        return float(v)
    except (ValueError, TypeError):
        raise HTTPException(400, f"'{field}' must be a number (got {v!r})")


def _excel_int(v, field):
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    try:
        return int(float(v))
    except (ValueError, TypeError):
        raise HTTPException(400, f"'{field}' must be a whole number (got {v!r})")


class InboundItemIn(BaseModel):
    sku_id: str
    qty: int
    location_id: str
    bag_color: Optional[str] = None
    batch_no: Optional[str] = None
    manufacture_date: Optional[str] = None
    expiry_date: Optional[str] = None
    barcode: Optional[str] = None


class InboundIn(BaseModel):
    po_number: str
    supplier: str
    vendor_id: Optional[str] = None
    expected_date: str
    items: List[InboundItemIn]


class PutawayScanIn(BaseModel):
    barcode: str


class PickScanIn(BaseModel):
    barcode: str


class FlowLaneAssignIn(BaseModel):
    sku_id: Optional[str] = None


class ZoneProvisionIn(BaseModel):
    zone_code: str
    zone_name: str
    temperature: Optional[float] = 22.0
    rows: List[dict]


class ZoneCreateIn(BaseModel):
    zone_code: str
    zone_name: str
    temperature: Optional[float] = 22.0


class ZoneEditIn(BaseModel):
    zone_name: str
    temperature: Optional[float] = None


class OutboundItemIn(BaseModel):
    sku_id: str
    qty: int


class OutboundIn(BaseModel):
    so_number: str
    customer: str
    customer_id: Optional[str] = None
    items: List[OutboundItemIn]


class StorageLaneInboundIn(BaseModel):
    level: int
    sku_id: str
    batch_no: Optional[str] = None
    manufacture_date: Optional[str] = None
    expiry_date: Optional[str] = None
    pallet_code: Optional[str] = None


class StorageLaneOutboundIn(BaseModel):
    level: int
    ref: Optional[str] = None


class ShuttleInboundIn(BaseModel):
    lane_no: int
    level_no: int
    sku_id: str
    batch_no: Optional[str] = None
    manufacture_date: Optional[str] = None
    expiry_date: Optional[str] = None
    pallet_code: Optional[str] = None


class ShuttleOutboundIn(BaseModel):
    lane_no: int
    level_no: int
    ref: Optional[str] = None


# ---------- Auth dep ----------
async def get_user(request: Request) -> dict:
    token = request.cookies.get("access_token")
    if not token:
        auth = request.headers.get("Authorization", "")
        if auth.startswith("Bearer "):
            token = auth[7:]
    if not token:
        raise HTTPException(401, "Not authenticated")
    try:
        payload = jwt.decode(token, jwt_secret(), algorithms=[JWT_ALGO])
    except jwt.ExpiredSignatureError:
        raise HTTPException(401, "Token expired")
    except jwt.InvalidTokenError:
        raise HTTPException(401, "Invalid token")
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        row = await conn.fetchrow(
            "SELECT id, email, name, role FROM users WHERE id = $1",
            payload["sub"],
        )
    if not row:
        raise HTTPException(401, "User not found")
    return dict(row)


def require_role(*roles: str):
    async def _checker(user: dict = Depends(get_user)):
        if user["role"] not in roles:
            raise HTTPException(403, "Forbidden")
        return user
    return _checker


def set_auth_cookie(resp: Response, token: str):
    resp.set_cookie(
        key="access_token",
        value=token,
        httponly=True,
        secure=True,
        samesite="none",
        max_age=60 * 60 * 24 * 7,
        path="/",
    )


# ─── inbound / outbound assembly helpers ─────────────────────────────────────

async def _assemble_inbound(conn, row) -> dict:
    order = dict(row)
    items = await conn.fetch(
        "SELECT * FROM inbound_items WHERE inbound_id = $1 ORDER BY id", order["id"]
    )
    order["items"] = [dict(it) for it in items]
    return order


async def _assemble_outbound(conn, row) -> dict:
    order = dict(row)
    items = await conn.fetch(
        "SELECT * FROM outbound_items WHERE outbound_id = $1", order["id"]
    )
    order["items"] = [dict(it) for it in items]
    return order


# ---------- Auth Routes ----------
@api.post("/auth/register")
async def register(body: RegisterIn, response: Response):
    email = body.email.lower()
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        existing = await conn.fetchrow("SELECT id FROM users WHERE email = $1", email)
        if existing:
            raise HTTPException(400, "Email already registered")
        uid = str(uuid.uuid4())
        await conn.execute(
            "INSERT INTO users (id, email, name, role, password_hash, created_at) VALUES ($1,$2,$3,$4,$5,$6)",
            uid, email, body.name, body.role, hash_pw(body.password), now_iso(),
        )
    token = create_token(uid, email)
    set_auth_cookie(response, token)
    return {"id": uid, "email": email, "name": body.name, "role": body.role, "token": token}


@api.post("/auth/login")
async def login(body: LoginIn, response: Response):
    email = body.email.lower()
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        row = await conn.fetchrow("SELECT * FROM users WHERE email = $1", email)
    if not row or not verify_pw(body.password, row["password_hash"]):
        raise HTTPException(401, "Invalid email or password")
    token = create_token(row["id"], email)
    set_auth_cookie(response, token)
    return {"id": row["id"], "email": email, "name": row["name"], "role": row["role"], "token": token}


@api.post("/auth/logout")
async def logout(response: Response):
    response.delete_cookie("access_token", path="/")
    return {"ok": True}


@api.get("/auth/me")
async def me(user: dict = Depends(get_user)):
    return {"id": user["id"], "email": user["email"], "name": user["name"], "role": user["role"]}


# ---------- Inventory ----------
@api.get("/inventory/skus")
async def list_skus(q: Optional[str] = None, category: Optional[str] = None, user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        if q and category:
            rows = await conn.fetch(
                "SELECT * FROM skus WHERE category = $1 AND (sku_code ILIKE $2 OR name ILIKE $2) ORDER BY sku_code",
                category, f"%{q}%",
            )
        elif q:
            rows = await conn.fetch(
                "SELECT * FROM skus WHERE sku_code ILIKE $1 OR name ILIKE $1 ORDER BY sku_code",
                f"%{q}%",
            )
        elif category:
            rows = await conn.fetch("SELECT * FROM skus WHERE category = $1 ORDER BY sku_code", category)
        else:
            rows = await conn.fetch("SELECT * FROM skus ORDER BY sku_code")
    return _db.rl(rows)


@api.get("/inventory/stock-by-color")
async def stock_by_color(user: dict = Depends(get_user)):
    """On-hand quantity grouped per SKU and actual bag color, used to build the
    'Inventory by Bag Color' breakdown in the Inventory Master."""
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        rows = await conn.fetch(
            "SELECT sku_id, bag_color, SUM(qty)::int AS on_hand "
            "FROM stock WHERE qty > 0 GROUP BY sku_id, bag_color"
        )
    return _db.rl(rows)


@api.post("/inventory/skus")
async def create_sku(body: SkuIn, user: dict = Depends(require_role("admin", "manager"))):
    _validate_sku_fields(body)
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        existing = await conn.fetchrow("SELECT id FROM skus WHERE sku_code = $1", body.sku_code)
        if existing:
            raise HTTPException(400, "SKU already exists")
        sid = str(uuid.uuid4())
        ts = now_iso()
        await conn.execute(
            "INSERT INTO skus (id, sku_code, name, category, bag_color, weight_per_bag, "
            "bags_per_pallet, dimensions, unit, unit_price, reorder_level, total_stock, created_at) "
            "VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10,$11,$12,$13)",
            sid, body.sku_code, body.name, body.category,
            body.bag_color, _to_decimal(body.weight_per_bag), body.bags_per_pallet, body.dimensions,
            body.unit, body.unit_price, body.reorder_level, 0, ts,
        )
    return {**body.model_dump(), "id": sid, "total_stock": 0, "created_at": ts}


@api.put("/inventory/skus/{sku_id}")
async def update_sku(sku_id: str, body: SkuIn, user: dict = Depends(require_role("admin", "manager"))):
    _validate_sku_fields(body)
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        res = await conn.execute(
            "UPDATE skus SET sku_code=$1, name=$2, category=$3, bag_color=$4, weight_per_bag=$5, "
            "bags_per_pallet=$6, dimensions=$7, unit=$8, unit_price=$9, reorder_level=$10 WHERE id=$11",
            body.sku_code, body.name, body.category,
            body.bag_color, _to_decimal(body.weight_per_bag), body.bags_per_pallet, body.dimensions,
            body.unit, body.unit_price, body.reorder_level, sku_id,
        )
        if res == "UPDATE 0":
            raise HTTPException(404, "SKU not found")
        row = await conn.fetchrow("SELECT * FROM skus WHERE id = $1", sku_id)
    return dict(row)


@api.delete("/inventory/skus/{sku_id}")
async def delete_sku(sku_id: str, user: dict = Depends(require_role("admin"))):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        await conn.execute("DELETE FROM stock WHERE sku_id = $1", sku_id)
        await conn.execute("DELETE FROM skus WHERE id = $1", sku_id)
    return {"ok": True}


@api.get("/inventory/skus/export")
async def export_skus(user: dict = Depends(get_user)):
    """Download the full SKU catalog as an .xlsx workbook. The same column
    layout can be re-uploaded via the import endpoint (the On Hand column is
    read-only and ignored on import)."""
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        rows = await conn.fetch("SELECT * FROM skus ORDER BY sku_code")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "SKUs"
    headers = [label for _, label in SKU_EXCEL_COLUMNS] + ["On Hand"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for r in rows:
        d = dict(r)
        line = []
        for key, _ in SKU_EXCEL_COLUMNS:
            v = d.get(key)
            if isinstance(v, Decimal):
                v = float(v)
            line.append(v)
        line.append(d.get("total_stock"))
        ws.append(line)

    for i, _ in enumerate(headers, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = 18

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"skus_{datetime.now(timezone.utc).strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


@api.post("/inventory/skus/import")
async def import_skus(
    file: UploadFile = File(...),
    user: dict = Depends(require_role("admin", "manager")),
):
    """Upsert SKUs from an uploaded .xlsx file, matching rows to existing SKUs
    by sku_code (create when new, update when found). Returns a per-row summary
    so the operator can see what was created/updated and which rows failed."""
    if not file.filename or not file.filename.lower().endswith((".xlsx", ".xlsm")):
        raise HTTPException(400, "Please upload an Excel .xlsx file")

    data = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    except Exception:
        raise HTTPException(400, "Could not read the file. Make sure it is a valid .xlsx workbook.")

    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        raise HTTPException(400, "The sheet is empty")

    # Accept either the friendly header label or the raw field name.
    label_to_key = {label.lower(): key for key, label in SKU_EXCEL_COLUMNS}
    for key, _ in SKU_EXCEL_COLUMNS:
        label_to_key[key.lower()] = key

    col_index = {}
    for idx, h in enumerate(header or []):
        if h is None:
            continue
        norm = str(h).strip().lower()
        if norm in label_to_key:
            col_index[label_to_key[norm]] = idx

    for required in ("sku_code", "name", "category"):
        if required not in col_index:
            label = dict(SKU_EXCEL_COLUMNS)[required]
            raise HTTPException(400, f"Missing required column: {label}")

    def cell_getter(raw):
        def cell(key):
            idx = col_index.get(key)
            if idx is None or idx >= len(raw):
                return None
            return raw[idx]
        return cell

    created = 0
    updated = 0
    errors = []
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        for n, raw in enumerate(rows_iter, start=2):
            if raw is None or all(c is None or str(c).strip() == "" for c in raw):
                continue
            cell = cell_getter(raw)
            try:
                payload = {
                    "sku_code": _excel_str(cell("sku_code")),
                    "name": _excel_str(cell("name")),
                    "category": _excel_str(cell("category")),
                    "bag_color": _excel_str(cell("bag_color")),
                    "dimensions": _excel_str(cell("dimensions")),
                    "unit": _excel_str(cell("unit")) or "EA",
                    "weight_per_bag": _excel_float(cell("weight_per_bag"), "Weight per Bag"),
                    "bags_per_pallet": _excel_int(cell("bags_per_pallet"), "Bags per Pallet"),
                    "unit_price": _excel_float(cell("unit_price"), "Unit Price") or 0.0,
                    "reorder_level": _excel_int(cell("reorder_level"), "Reorder Level"),
                }
                if not payload["sku_code"]:
                    raise HTTPException(400, "SKU Code is required")
                if not payload["name"]:
                    raise HTTPException(400, "Name is required")
                if not payload["category"]:
                    raise HTTPException(400, "Category is required")
                if payload["reorder_level"] is None:
                    payload["reorder_level"] = 10

                body = SkuIn(**payload)
                _validate_sku_fields(body)

                existing = await conn.fetchrow("SELECT id FROM skus WHERE sku_code = $1", body.sku_code)
                if existing:
                    await conn.execute(
                        "UPDATE skus SET name=$1, category=$2, bag_color=$3, weight_per_bag=$4, "
                        "bags_per_pallet=$5, dimensions=$6, unit=$7, unit_price=$8, reorder_level=$9 WHERE id=$10",
                        body.name, body.category, body.bag_color, _to_decimal(body.weight_per_bag),
                        body.bags_per_pallet, body.dimensions, body.unit, body.unit_price,
                        body.reorder_level, existing["id"],
                    )
                    updated += 1
                else:
                    sid = str(uuid.uuid4())
                    ts = now_iso()
                    await conn.execute(
                        "INSERT INTO skus (id, sku_code, name, category, bag_color, weight_per_bag, "
                        "bags_per_pallet, dimensions, unit, unit_price, reorder_level, total_stock, created_at) "
                        "VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10,$11,$12,$13)",
                        sid, body.sku_code, body.name, body.category, body.bag_color,
                        _to_decimal(body.weight_per_bag), body.bags_per_pallet, body.dimensions,
                        body.unit, body.unit_price, body.reorder_level, 0, ts,
                    )
                    created += 1
            except HTTPException as e:
                errors.append({"row": n, "sku_code": _excel_str(cell("sku_code")) or "", "error": e.detail})
            except Exception as e:
                errors.append({"row": n, "sku_code": _excel_str(cell("sku_code")) or "", "error": str(e)})

    return {
        "created": created,
        "updated": updated,
        "errors": errors,
        "total": created + updated + len(errors),
    }


@api.get("/inventory/categories")
async def categories(user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        rows = await conn.fetch("SELECT DISTINCT category FROM skus ORDER BY category")
    return [r["category"] for r in rows]


@api.get("/inventory/skus/{sku_id}/stock")
async def sku_stock(sku_id: str, user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        rows = await conn.fetch(
            "SELECT s.qty, l.* FROM stock s JOIN locations l ON l.id = s.location_id "
            "WHERE s.sku_id = $1 AND s.qty > 0",
            sku_id,
        )
    out = []
    for r in rows:
        d = dict(r)
        out.append({"location": {k: d[k] for k in d if k != "qty"}, "qty": d["qty"]})
    return out


@api.get("/inventory/skus/{sku_id}/pallets")
async def sku_pallets(sku_id: str, user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        rows = await conn.fetch(
            "SELECT s.qty, s.batch_no, s.expiry_date, s.manufacture_date, s.received_date, "
            "l.id AS loc_id, l.code AS loc_code, l.zone "
            "FROM stock s JOIN locations l ON l.id = s.location_id "
            "WHERE s.sku_id = $1 AND s.qty > 0",
            sku_id,
        )
    out = [
        {
            "location": {"id": r["loc_id"], "code": r["loc_code"], "zone": r["zone"]},
            "qty": r["qty"],
            "batch_no": r["batch_no"],
            "expiry_date": r["expiry_date"],
            "manufacture_date": r["manufacture_date"],
            "received_date": r["received_date"],
        }
        for r in rows
    ]
    out.sort(key=lambda x: (x.get("expiry_date") is None, x.get("expiry_date") or "9999-99-99"))
    return out


# ---------- Storage / Locations ----------
@api.get("/storage/locations")
async def list_locations(zone: Optional[str] = None, user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        if zone:
            rows = await conn.fetch(
                "SELECT * FROM locations WHERE zone = $1 ORDER BY zone, code", zone
            )
        else:
            rows = await conn.fetch("SELECT * FROM locations ORDER BY zone, code")
    return _db.rl(rows)


@api.get("/storage/zones")
async def zones(user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        rows = await conn.fetch(
            "SELECT zone, zone_name, temperature, "
            "SUM(capacity) AS capacity, "
            "SUM(CASE WHEN occupied > 0 THEN 1 ELSE 0 END) AS occupied, "
            "COUNT(*) AS bins "
            "FROM locations GROUP BY zone, zone_name, temperature ORDER BY zone"
        )
        placeholders = await conn.fetch("SELECT * FROM zones_meta ORDER BY zone")
    out = []
    zone_set = set()
    for r in rows:
        out.append({
            "zone": r["zone"],
            "name": r["zone_name"] or r["zone"],
            "temperature": r["temperature"],
            "capacity": r["capacity"],
            "occupied": r["occupied"],
            "bins": r["bins"],
        })
        zone_set.add(r["zone"])
    for p in placeholders:
        if p["zone"] not in zone_set:
            out.append({
                "zone": p["zone"],
                "name": p["name"],
                "temperature": p["temperature"],
                "capacity": 0,
                "occupied": 0,
                "bins": 0,
                "placeholder": True,
            })
    return out


@api.get("/storage/lanes/{row}/{lane_number}/contents")
async def lane_contents(row: str, lane_number: int, user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        bins = await conn.fetch(
            "SELECT * FROM locations WHERE row_label = $1 AND lane_number = $2 "
            "ORDER BY level, position",
            row, lane_number,
        )
        bin_ids = [b["id"] for b in bins]
        if not bin_ids:
            return []
        stock_rows = await conn.fetch(
            "SELECT s.*, sk.id AS sk_id, sk.sku_code, sk.name AS sk_name, sk.category "
            "FROM stock s JOIN skus sk ON sk.id = s.sku_id "
            "WHERE s.location_id = ANY($1::text[]) AND s.qty > 0",
            bin_ids,
        )
    stock_map = {}
    for s in stock_rows:
        stock_map[s["location_id"]] = {
            "sku": {"id": s["sk_id"], "sku_code": s["sku_code"], "name": s["sk_name"], "category": s["category"]},
            "qty": s["qty"],
            "batch_no": s["batch_no"],
            "manufacture_date": s["manufacture_date"],
            "expiry_date": s["expiry_date"],
            "received_date": s["received_date"],
            "pallet_code": s["pallet_code"],
        }
    return [{"bin": dict(b), "item": stock_map.get(b["id"])} for b in bins]


@api.get("/storage/bins/{bin_id}/pallet")
async def bin_pallet_detail(bin_id: str, user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        loc = await conn.fetchrow("SELECT * FROM locations WHERE id = $1", bin_id)
        if not loc:
            raise HTTPException(404, "Bin not found")
        stock = await conn.fetchrow(
            "SELECT s.*, sk.id AS sk_id, sk.sku_code, sk.name AS sk_name, sk.category, sk.unit "
            "FROM stock s JOIN skus sk ON sk.id = s.sku_id "
            "WHERE s.location_id = $1 AND s.qty > 0",
            bin_id,
        )
    loc_d = dict(loc)
    item = None
    if stock:
        item = {
            "sku": {"id": stock["sk_id"], "sku_code": stock["sku_code"], "name": stock["sk_name"],
                    "category": stock["category"], "unit": stock["unit"]},
            "qty": stock["qty"],
            "batch_no": stock["batch_no"],
            "pallet_code": stock["pallet_code"],
            "manufacture_date": stock["manufacture_date"],
            "expiry_date": stock["expiry_date"],
            "received_date": stock["received_date"],
        }
    return {
        "bin": {k: loc_d[k] for k in ("id", "code", "level", "position", "zone", "row_label",
                                        "lane_number", "weight_capacity_kg", "rack_type")},
        "item": item,
    }


# ---------- Storage Lane Inbound / Outbound (drive-in racks) ----------
@api.post("/storage/lanes/{zone}/{row}/{lane_number}/inbound")
async def storage_lane_inbound(
    zone: str, row: str, lane_number: int,
    body: StorageLaneInboundIn,
    user: dict = Depends(require_role("admin", "manager", "operator")),
):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        bins = await conn.fetch(
            "SELECT * FROM locations WHERE zone=$1 AND row_label=$2 AND lane_number=$3 AND level=$4 ORDER BY position",
            zone, row, lane_number, body.level,
        )
        if not bins:
            raise HTTPException(404, f"Lane {row}-{lane_number} level {body.level} not found in zone {zone}")
        empty = [b for b in bins if b["occupied"] == 0]
        if not empty:
            raise HTTPException(409, f"Lane level {body.level} is fully occupied")
        target = max(empty, key=lambda b: b["position"])
        pallet_code = body.pallet_code or f"PLT-{zone}-{row}{lane_number:02d}L{body.level:02d}P{target['position']:02d}"
        ts = now_iso()
        await conn.execute(
            "INSERT INTO stock (id, sku_id, location_id, qty, batch_no, manufacture_date, expiry_date, received_date, pallet_code, pallet_status, original_qty) "
            "VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10,$11)",
            str(uuid.uuid4()), body.sku_id, target["id"], 1,
            body.batch_no, body.manufacture_date, body.expiry_date, ts, pallet_code, "full", 1,
        )
        await conn.execute("UPDATE locations SET occupied = occupied + 1 WHERE id = $1", target["id"])
    return {
        "ok": True, "bin_code": target["code"],
        "placed_at_level": body.level, "placed_at_position": target["position"],
        "pallet_code": pallet_code,
    }


@api.post("/storage/lanes/{zone}/{row}/{lane_number}/outbound")
async def storage_lane_outbound(
    zone: str, row: str, lane_number: int,
    body: StorageLaneOutboundIn,
    user: dict = Depends(require_role("admin", "manager", "operator")),
):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        bins = await conn.fetch(
            "SELECT * FROM locations WHERE zone=$1 AND row_label=$2 AND lane_number=$3 AND level=$4 AND occupied > 0 "
            "ORDER BY position",
            zone, row, lane_number, body.level,
        )
        if not bins:
            raise HTTPException(409, f"No pallets in level {body.level} of lane {row}-{lane_number}")
        target = bins[0]
        stock = await conn.fetchrow(
            "SELECT s.*, sk.sku_code, sk.name AS sku_name FROM stock s JOIN skus sk ON sk.id = s.sku_id "
            "WHERE s.location_id = $1 AND s.qty > 0", target["id"]
        )
        if not stock:
            raise HTTPException(500, "Location shows occupied but no stock record")
        dispatched = {
            "pallet_code": stock["pallet_code"], "sku_code": stock["sku_code"],
            "sku_name": stock["sku_name"], "batch_no": stock["batch_no"],
            "expiry_date": stock["expiry_date"], "bin_code": target["code"],
            "level": body.level, "position": target["position"],
        }
        await conn.execute("DELETE FROM stock WHERE location_id = $1 AND qty > 0", target["id"])
        await conn.execute("UPDATE locations SET occupied = 0 WHERE id = $1", target["id"])
    return {"ok": True, "dispatched": dispatched}


# ---------- Blueprint / Zone Provision ----------
@api.post("/storage/parse-blueprint")
async def parse_blueprint(
    file: UploadFile = File(...),
    user: dict = Depends(require_role("admin", "manager")),
):
    import json as json_lib
    import base64
    import fitz

    SYSTEM_PROMPT = (
        "You are a warehouse blueprint analyst. Extract rack/lane configuration from the blueprint "
        "and return STRICT JSON only (no prose, no markdown). Schema: "
        '{"zone_name": "string", "rows": ['
        '{"row": "A|B|C|...", "rack_type": "A|B|C", "lanes": int, "lane_start": int, '
        '"levels": int, "depth": int, "weight_kg": int}]}. '
        "If a field is unclear, infer reasonable defaults (levels=4, depth=4, weight_kg=8000)."
    )

    raw = await file.read()
    try:
        doc = fitz.open(stream=raw, filetype="pdf")
    except Exception as e:
        raise HTTPException(400, f"Could not read PDF: {e}")

    # --- Always generate a thumbnail for reference in the edit step ---
    try:
        first_page = doc[0]
        page_w = first_page.rect.width
        thumb_scale = 1200 / page_w if page_w > 0 else 0.2
        pix = first_page.get_pixmap(matrix=fitz.Matrix(thumb_scale, thumb_scale))
        thumb_b64 = base64.standard_b64encode(pix.tobytes("jpeg", jpg_quality=75)).decode()
        blueprint_image = f"data:image/jpeg;base64,{thumb_b64}"
    except Exception:
        blueprint_image = None

    # --- Try text extraction first (vector/native PDFs) ---
    text = ""
    for page in doc:
        text += page.get_text() + "\n"

    try:
        _client, _model = _get_anthropic_client()

        if text.strip():
            # Text-based PDF: send as plain text
            _resp = _client.messages.create(
                model=_model,
                max_tokens=4096,
                system=SYSTEM_PROMPT,
                messages=[{"role": "user", "content": f"BLUEPRINT TEXT:\n\n{text[:30000]}"}],
            )
        else:
            # Image-based (scanned) PDF: render pages to PNG and use Claude vision
            images_content = []
            max_pages = min(len(doc), 5)
            for i in range(max_pages):
                page = doc[i]
                # Cap longest side at 3500px to stay within Claude's image size limit
                rect = doc[i].rect
                max_side = max(rect.width, rect.height)
                ai_scale = min(2.0, 3500 / max_side) if max_side > 0 else 1.0
                pix_ai = page.get_pixmap(matrix=fitz.Matrix(ai_scale, ai_scale))
                png_bytes = pix_ai.tobytes("png")
                b64 = base64.standard_b64encode(png_bytes).decode()
                images_content.append({
                    "type": "image",
                    "source": {"type": "base64", "media_type": "image/png", "data": b64},
                })
            images_content.append({
                "type": "text",
                "text": (
                    "Extract the COMPLETE warehouse rack/lane configuration from ALL rows visible "
                    "in these blueprint images. Do not stop early — include every row you can see."
                ),
            })
            _resp = _client.messages.create(
                model=_model,
                max_tokens=4096,
                system=SYSTEM_PROMPT,
                messages=[{"role": "user", "content": images_content}],
            )

        doc.close()
        response = _resp.content[0].text
        cleaned = response.strip()
        if cleaned.startswith("```"):
            cleaned = cleaned.split("```")[1]
            if cleaned.startswith("json"):
                cleaned = cleaned[4:]
        config = json_lib.loads(cleaned.strip())
        return {"config": config, "raw_response": response, "blueprint_image": blueprint_image}
    except Exception as e:
        doc.close()
        return {
            "config": {
                "zone_name": "New Zone",
                "rows": [{"row": "A", "rack_type": "A", "lanes": 5, "lane_start": 1,
                          "levels": 4, "depth": 4, "weight_kg": 8000}],
            },
            "warning": f"AI parse failed, returning defaults: {str(e)[:200]}",
            "blueprint_image": blueprint_image,
        }


@api.post("/storage/zones/{zone_code}/provision")
async def provision_zone(
    zone_code: str, body: ZoneProvisionIn,
    user: dict = Depends(require_role("admin", "manager")),
):
    if zone_code != body.zone_code:
        raise HTTPException(400, "Zone code mismatch")
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        cnt = await conn.fetchval("SELECT COUNT(*) FROM locations WHERE zone = $1", zone_code)
        if cnt > 0:
            raise HTTPException(400, "Zone already provisioned")
        locs = []
        for row_cfg in body.rows:
            row_label = row_cfg["row"]
            rack_type = row_cfg.get("rack_type", "A")
            lanes_count = int(row_cfg["lanes"])
            lane_start = int(row_cfg.get("lane_start", 1))
            levels = int(row_cfg.get("levels", 4))
            depth = int(row_cfg.get("depth", 4))
            weight_kg = int(row_cfg.get("weight_kg", 8000))
            for i in range(lanes_count):
                lane_num = lane_start + i
                for level in range(1, levels + 1):
                    for pos in range(1, depth + 1):
                        locs.append((
                            str(uuid.uuid4()),
                            f"{zone_code}-{row_label}-L{lane_num:02d}-LV{level}-P{pos:02d}",
                            zone_code, body.zone_name, body.temperature,
                            rack_type, row_label, lane_num, level, pos, depth, levels, weight_kg, 1, 0,
                        ))
        if locs:
            await conn.executemany(
                "INSERT INTO locations (id, code, zone, zone_name, temperature, rack_type, row_label, "
                "lane_number, level, position, depth, levels, weight_capacity_kg, capacity, occupied) "
                "VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10,$11,$12,$13,$14,$15)",
                locs,
            )
        await conn.execute(
            "INSERT INTO zones_meta (zone, name, temperature, placeholder) VALUES ($1,$2,$3,$4) "
            "ON CONFLICT (zone) DO UPDATE SET name=$2, temperature=$3, placeholder=$4",
            zone_code, body.zone_name, body.temperature, False,
        )
    return {"ok": True, "bins_created": len(locs)}


@api.post("/storage/zones")
async def create_zone(body: ZoneCreateIn, user: dict = Depends(require_role("admin", "manager"))):
    code = body.zone_code.strip().upper()
    if not code:
        raise HTTPException(400, "Zone code required")
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        existing = await conn.fetchrow("SELECT zone FROM zones_meta WHERE zone = $1", code)
        if existing:
            raise HTTPException(400, f"Zone {code} already exists")
        await conn.execute(
            "INSERT INTO zones_meta (zone, name, temperature, placeholder) VALUES ($1,$2,$3,$4)",
            code, body.zone_name.strip() or code, body.temperature, True,
        )
    return {"ok": True, "zone": code}


@api.put("/storage/zones/{zone_code}")
async def edit_zone(zone_code: str, body: ZoneEditIn, user: dict = Depends(require_role("admin", "manager"))):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        if body.temperature is not None:
            await conn.execute(
                "INSERT INTO zones_meta (zone, name, temperature, placeholder) VALUES ($1,$2,$3,FALSE) "
                "ON CONFLICT (zone) DO UPDATE SET name=$2, temperature=$3",
                zone_code, body.zone_name.strip(), body.temperature,
            )
            await conn.execute(
                "UPDATE locations SET zone_name=$1, temperature=$2 WHERE zone=$3",
                body.zone_name.strip(), body.temperature, zone_code,
            )
        else:
            await conn.execute(
                "INSERT INTO zones_meta (zone, name, placeholder) VALUES ($1,$2,FALSE) "
                "ON CONFLICT (zone) DO UPDATE SET name=$2",
                zone_code, body.zone_name.strip(),
            )
            await conn.execute(
                "UPDATE locations SET zone_name=$1 WHERE zone=$2",
                body.zone_name.strip(), zone_code,
            )
    return {"ok": True}


@api.delete("/storage/zones/{zone_code}")
async def delete_zone(zone_code: str, user: dict = Depends(require_role("admin"))):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        occupied = await conn.fetchval(
            "SELECT COUNT(*) FROM locations WHERE zone=$1 AND occupied > 0", zone_code
        )
        if occupied > 0:
            raise HTTPException(409, f"Zone {zone_code} still has {occupied} occupied slot(s). Clear stock before deleting.")
        deleted = await conn.fetchval("SELECT COUNT(*) FROM locations WHERE zone=$1", zone_code)
        await conn.execute("DELETE FROM locations WHERE zone=$1", zone_code)
        await conn.execute("DELETE FROM zones_meta WHERE zone=$1", zone_code)
    return {"ok": True, "bins_deleted": deleted}


# ---------- Storage Lanes ----------
@api.get("/storage/lanes")
async def lanes(zone: Optional[str] = None, user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        if zone:
            locs = await conn.fetch("SELECT * FROM locations WHERE zone=$1", zone)
        else:
            locs = await conn.fetch("SELECT * FROM locations")
    lanes_map: dict = {}
    for l in locs:
        row_label = l["row_label"] or "?"
        lane_num = l["lane_number"] or 0
        key = (row_label, lane_num)
        lane = lanes_map.setdefault(key, {
            "row": row_label, "lane_number": lane_num,
            "rack_type": l["rack_type"], "levels": l["levels"] or 4,
            "depth": l["depth"] or 0, "weight_capacity_kg": l["weight_capacity_kg"] or 0,
            "sku_assignment": l["sku_assignment"], "bins": [],
        })
        lane["bins"].append({
            "id": l["id"], "code": l["code"], "level": l["level"],
            "position": l["position"], "occupied": l["occupied"], "capacity": l["capacity"],
        })
    out = []
    for lane in lanes_map.values():
        lane["bins"].sort(key=lambda b: (b.get("level") or 0, b.get("position") or 0))
        lane["total_slots"] = len(lane["bins"])
        lane["filled_slots"] = sum(1 for b in lane["bins"] if (b["occupied"] or 0) > 0)
        out.append(lane)
    out.sort(key=lambda x: (x["row"], x["lane_number"]))
    return out


@api.get("/storage/flow-lanes")
async def flow_lanes_summary(user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        locs = await conn.fetch("SELECT * FROM locations WHERE rack_type = 'flow_rack'")
    lanes_map: dict = {}
    for l in locs:
        key = (l["zone"], l["row_label"] or "?", l["lane_number"] or 0)
        lane = lanes_map.setdefault(key, {
            "zone": l["zone"], "zone_name": l["zone_name"] or l["zone"],
            "row": l["row_label"] or "?", "lane_number": l["lane_number"] or 0,
            "depth": l["depth"] or 0, "levels": l["levels"] or 1,
            "weight_capacity_kg": l["weight_capacity_kg"] or 0,
            "sku_assignment": l["sku_assignment"],
            "total_slots": 0, "filled_slots": 0,
        })
        lane["total_slots"] += 1
        if (l["occupied"] or 0) > 0:
            lane["filled_slots"] += 1
    pool2 = await _db.get_pool()
    out = list(lanes_map.values())
    for lane in out:
        if lane.get("sku_assignment"):
            async with pool2.acquire() as conn:
                sku = await conn.fetchrow(
                    "SELECT sku_code, name, unit FROM skus WHERE id=$1", lane["sku_assignment"]
                )
            lane["sku"] = dict(sku) if sku else None
        else:
            lane["sku"] = None
    out.sort(key=lambda x: (x["zone"], x["row"], x["lane_number"]))
    return out


@api.patch("/storage/lanes/{zone}/{row}/{lane_number}/assign-sku")
async def assign_flow_lane_sku(
    zone: str, row: str, lane_number: int,
    body: FlowLaneAssignIn,
    user: dict = Depends(require_role("admin", "manager")),
):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        res = await conn.execute(
            "UPDATE locations SET sku_assignment=$1 WHERE zone=$2 AND row_label=$3 AND lane_number=$4 AND rack_type='flow_rack'",
            body.sku_id, zone, row, lane_number,
        )
    if res == "UPDATE 0":
        raise HTTPException(404, "Flow rack lane not found")
    return {"ok": True}


# ════════════════════════════════════════════════════
# SHUTTLE FIFO DEEP LANE — SHUTTLE_ZONE_A
# ════════════════════════════════════════════════════

@api.get("/shuttle/summary")
async def shuttle_summary(user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        total = await conn.fetchval("SELECT COUNT(*) FROM locations WHERE zone=$1", SHUTTLE_ZONE)
        occupied = await conn.fetchval(
            "SELECT COUNT(*) FROM locations WHERE zone=$1 AND occupied > 0", SHUTTLE_ZONE
        )
    return {
        "zone": SHUTTLE_ZONE, "zone_name": "Shuttle FIFO — Cold Storage",
        "total_bins": total, "occupied_bins": occupied,
        "empty_bins": total - occupied,
        "utilization_pct": round(occupied / total * 100, 1) if total else 0,
        "lanes": 7, "levels": 5, "depth": 50,
    }


@api.get("/shuttle/lanes")
async def shuttle_lanes(user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        rows = await conn.fetch(
            "SELECT lane_number, level, COUNT(*) AS total, SUM(CASE WHEN occupied>0 THEN 1 ELSE 0 END) AS occ, "
            "MAX(CASE WHEN occupied>0 THEN position ELSE 0 END) AS last_depth "
            "FROM locations WHERE zone=$1 GROUP BY lane_number, level ORDER BY lane_number, level",
            SHUTTLE_ZONE,
        )
    out = []
    for r in rows:
        total = r["total"]
        occ = r["occ"]
        out.append({
            "lane_no": r["lane_number"], "level_no": r["level"],
            "total": total, "occupied": occ, "empty": total - occ,
            "last_inbound_depth": r["last_depth"] or 0,
            "utilization_pct": round(occ / total * 100, 1) if total else 0,
        })
    return out


@api.get("/shuttle/lanes/{lane_no}/{level_no}")
async def shuttle_lane_detail(lane_no: int, level_no: int, user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        locs = await conn.fetch(
            "SELECT l.*, s.id AS s_id, s.sku_id, s.batch_no, s.expiry_date, s.manufacture_date, "
            "s.received_date, s.pallet_code, sk.sku_code, sk.name AS sk_name "
            "FROM locations l "
            "LEFT JOIN stock s ON s.location_id = l.id AND s.qty > 0 "
            "LEFT JOIN skus sk ON sk.id = s.sku_id "
            "WHERE l.zone=$1 AND l.lane_number=$2 AND l.level=$3 ORDER BY l.position",
            SHUTTLE_ZONE, lane_no, level_no,
        )
    out = []
    for l in locs:
        stock = None
        if l["s_id"]:
            stock = {
                "sku_id": l["sku_id"], "sku_code": l["sku_code"] or "?",
                "sku_name": l["sk_name"] or "?", "batch_no": l["batch_no"],
                "expiry_date": l["expiry_date"], "manufacture_date": l["manufacture_date"],
                "received_date": l["received_date"], "pallet_code": l["pallet_code"],
            }
        out.append({
            "depth": l["position"], "code": l["code"],
            "bin_id": l["id"], "occupied": (l["occupied"] or 0) > 0, "stock": stock,
        })
    return out


@api.post("/shuttle/inbound")
async def shuttle_inbound(body: ShuttleInboundIn, user: dict = Depends(require_role("admin", "manager", "operator"))):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        locs = await conn.fetch(
            "SELECT * FROM locations WHERE zone=$1 AND lane_number=$2 AND level=$3 ORDER BY position",
            SHUTTLE_ZONE, body.lane_no, body.level_no,
        )
        if not locs:
            raise HTTPException(404, f"Lane {body.lane_no} Level {body.level_no} not found")
        occupied_positions = [l["position"] for l in locs if (l["occupied"] or 0) > 0]
        if len(occupied_positions) >= 50:
            raise HTTPException(409, "Lane is fully occupied (50/50)")
        next_depth = (max(occupied_positions) + 1) if occupied_positions else 1
        target_bin = next(l for l in locs if l["position"] == next_depth)
        ts = now_iso()
        pallet_code = body.pallet_code or f"PLT-SZA-{body.lane_no:02d}{body.level_no:02d}{next_depth:02d}"
        stock_id = str(uuid.uuid4())
        await conn.execute(
            "INSERT INTO stock (id, sku_id, location_id, qty, batch_no, manufacture_date, expiry_date, received_date, pallet_code, pallet_status, original_qty) "
            "VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10,$11)",
            stock_id, body.sku_id, target_bin["id"], 1,
            body.batch_no, body.manufacture_date, body.expiry_date, ts, pallet_code, "full", 1,
        )
        await conn.execute("UPDATE locations SET occupied = occupied + 1 WHERE id=$1", target_bin["id"])
        sku = await conn.fetchrow("SELECT sku_code FROM skus WHERE id=$1", body.sku_id)
        await conn.execute(
            "INSERT INTO shuttle_movements (id, lane_no, level_no, from_depth, to_depth, pallet_code, sku_code, movement_type, timestamp) "
            "VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9)",
            str(uuid.uuid4()), body.lane_no, body.level_no, None, next_depth,
            pallet_code, sku["sku_code"] if sku else "?", "inbound", ts,
        )
    return {
        "ok": True, "lane_no": body.lane_no, "level_no": body.level_no,
        "placed_at_depth": next_depth, "bin_code": target_bin["code"], "pallet_code": pallet_code,
    }


@api.post("/shuttle/outbound")
async def shuttle_outbound(body: ShuttleOutboundIn, user: dict = Depends(require_role("admin", "manager", "operator"))):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        async with conn.transaction():
            locs = await conn.fetch(
                "SELECT * FROM locations WHERE zone=$1 AND lane_number=$2 AND level=$3 ORDER BY position FOR UPDATE",
                SHUTTLE_ZONE, body.lane_no, body.level_no,
            )
            if not locs:
                raise HTTPException(404, f"Lane {body.lane_no} Level {body.level_no} not found")
            loc_by_depth = {l["position"]: l for l in locs}
            d01 = loc_by_depth.get(1)
            if not d01 or (d01["occupied"] or 0) == 0:
                raise HTTPException(409, "D01 is empty — FIFO requires picking from D01")
            d01_stock = await conn.fetchrow(
                "SELECT s.*, sk.sku_code, sk.name AS sku_name FROM stock s JOIN skus sk ON sk.id = s.sku_id "
                "WHERE s.location_id=$1 AND s.qty > 0", d01["id"]
            )
            if not d01_stock:
                raise HTTPException(500, "D01 shows occupied but stock record missing")
            dispatched = {
                "pallet_code": d01_stock["pallet_code"], "sku_code": d01_stock["sku_code"],
                "sku_name": d01_stock["sku_name"], "batch_no": d01_stock["batch_no"],
                "expiry_date": d01_stock["expiry_date"], "received_date": d01_stock["received_date"],
                "bin_code": d01["code"],
            }
            await conn.execute("DELETE FROM stock WHERE location_id=$1 AND qty > 0", d01["id"])
            await conn.execute("UPDATE locations SET occupied=0 WHERE id=$1", d01["id"])
            ts = now_iso()
            ref = body.ref or f"SO-SHUTTLE-{body.lane_no:02d}"
            await conn.execute(
                "INSERT INTO shuttle_movements (id, lane_no, level_no, from_depth, to_depth, pallet_code, sku_code, movement_type, ref, timestamp) "
                "VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10)",
                str(uuid.uuid4()), body.lane_no, body.level_no, 1, None,
                dispatched["pallet_code"], dispatched["sku_code"], "outbound", ref, ts,
            )
            # Shift: D2→D1, D3→D2, …
            shifts = 0
            for depth in range(2, 51):
                src = loc_by_depth.get(depth)
                if not src or (src["occupied"] or 0) == 0:
                    break
                dst = loc_by_depth.get(depth - 1)
                await conn.execute(
                    "UPDATE stock SET location_id=$1 WHERE location_id=$2 AND qty > 0",
                    dst["id"], src["id"],
                )
                await conn.execute("UPDATE locations SET occupied=1 WHERE id=$1", dst["id"])
                await conn.execute("UPDATE locations SET occupied=0 WHERE id=$1", src["id"])
                await conn.execute(
                    "INSERT INTO shuttle_movements (id, lane_no, level_no, from_depth, to_depth, sku_code, movement_type, ref, timestamp) "
                    "VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9)",
                    str(uuid.uuid4()), body.lane_no, body.level_no,
                    depth, depth - 1, dispatched["sku_code"], "shift", ref, ts,
                )
                shifts += 1
            await conn.execute(
                "UPDATE skus SET total_stock = total_stock - 1 WHERE id=$1", d01_stock["sku_id"]
            )
    return {
        "ok": True, "dispatched": dispatched,
        "pallets_shifted": shifts, "lane_no": body.lane_no, "level_no": body.level_no,
    }


@api.get("/shuttle/movements")
async def shuttle_movement_history(
    lane_no: Optional[int] = None,
    limit: int = 100,
    user: dict = Depends(get_user),
):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        if lane_no:
            rows = await conn.fetch(
                "SELECT * FROM shuttle_movements WHERE movement_type != 'shift' AND lane_no=$1 "
                "ORDER BY timestamp DESC LIMIT $2", lane_no, limit,
            )
        else:
            rows = await conn.fetch(
                "SELECT * FROM shuttle_movements WHERE movement_type != 'shift' "
                "ORDER BY timestamp DESC LIMIT $1", limit,
            )
    return _db.rl(rows)


# ---------- Transfer ----------

class TransferIn(BaseModel):
    stock_id: str
    to_location_id: str
    notes: Optional[str] = None


@api.get("/storage/pallets")
async def list_pallets(user: dict = Depends(get_user)):
    """All stock items currently placed in storage, enriched with SKU and location info."""
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        rows = await conn.fetch(
            "SELECT s.id, s.sku_id, s.qty, s.batch_no, s.expiry_date, s.received_date, "
            "s.pallet_code, s.bag_color, s.location_id, "
            "sk.sku_code, sk.name AS sku_name, sk.category, "
            "l.code AS location_code, l.zone, l.zone_name "
            "FROM stock s "
            "JOIN skus sk ON sk.id = s.sku_id "
            "JOIN locations l ON l.id = s.location_id "
            "WHERE s.qty > 0 "
            "ORDER BY l.zone, l.code, sk.sku_code"
        )
    return _db.rl(rows)


@api.get("/storage/bin-stock")
async def bin_stock(location_id: str, user: dict = Depends(get_user)):
    """Return all stock items currently occupying a specific location bin."""
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        rows = await conn.fetch(
            "SELECT s.id, s.qty, s.batch_no, s.expiry_date, s.received_date, "
            "s.pallet_code, s.bag_color, s.pallet_status, "
            "sk.sku_code, sk.name AS sku_name, sk.unit, sk.category "
            "FROM stock s "
            "JOIN skus sk ON sk.id = s.sku_id "
            "WHERE s.location_id = $1 AND s.qty > 0 "
            "ORDER BY s.received_date NULLS LAST",
            location_id,
        )
    return _db.rl(rows)


@api.post("/storage/transfer")
async def transfer_pallet(body: TransferIn, user: dict = Depends(require_role("admin", "manager"))):
    """Move a pallet (stock record) from its current location to a different location.
    Validates capacity, updates occupied counts, and writes an audit record."""
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        stock = await conn.fetchrow(
            "SELECT s.*, sk.sku_code, sk.name AS sku_name, "
            "l.code AS from_code, l.id AS from_loc_id "
            "FROM stock s "
            "JOIN skus sk ON sk.id = s.sku_id "
            "JOIN locations l ON l.id = s.location_id "
            "WHERE s.id = $1 AND s.qty > 0",
            body.stock_id,
        )
        if not stock:
            raise HTTPException(404, "Pallet not found or has zero quantity")
        if stock["location_id"] == body.to_location_id:
            raise HTTPException(400, "Source and destination are the same location")

        dest = await conn.fetchrow("SELECT * FROM locations WHERE id = $1", body.to_location_id)
        if not dest:
            raise HTTPException(404, "Destination location not found")
        if dest["occupied"] >= dest["capacity"]:
            raise HTTPException(
                409,
                f"Destination {dest['code']} is at full capacity ({dest['capacity']}/{dest['capacity']})",
            )

        ts = now_iso()
        tid = str(uuid.uuid4())

        await conn.execute("UPDATE stock SET location_id = $1 WHERE id = $2", body.to_location_id, body.stock_id)
        await conn.execute("UPDATE locations SET occupied = GREATEST(occupied - 1, 0) WHERE id = $1", stock["from_loc_id"])
        await conn.execute("UPDATE locations SET occupied = occupied + 1 WHERE id = $1", body.to_location_id)
        await conn.execute(
            "INSERT INTO transfers (id, stock_id, sku_id, sku_code, sku_name, pallet_code, qty, bag_color, "
            "from_location_id, from_code, to_location_id, to_code, notes, transferred_by, transferred_at, status) "
            "VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10,$11,$12,$13,$14,$15,$16)",
            tid, body.stock_id, stock["sku_id"], stock["sku_code"], stock["sku_name"],
            stock["pallet_code"], stock["qty"], stock["bag_color"],
            stock["from_loc_id"], stock["from_code"], body.to_location_id, dest["code"],
            body.notes, user["id"], ts, "initiated",
        )

    return {
        "id": tid,
        "pallet_code": stock["pallet_code"],
        "sku_code": stock["sku_code"],
        "sku_name": stock["sku_name"],
        "qty": stock["qty"],
        "bag_color": stock["bag_color"],
        "from_code": stock["from_code"],
        "to_code": dest["code"],
        "transferred_at": ts,
    }


@api.get("/storage/transfers")
async def list_transfers(limit: int = 50, user: dict = Depends(get_user)):
    """Recent pallet transfer history, newest first."""
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        rows = await conn.fetch(
            "SELECT t.*, u.name AS transferred_by_name, "
            "cu.name AS confirmed_by_name "
            "FROM transfers t "
            "LEFT JOIN users u ON u.id = t.transferred_by "
            "LEFT JOIN users cu ON cu.id = t.confirmed_by "
            "ORDER BY t.transferred_at DESC LIMIT $1",
            limit,
        )
    return _db.rl(rows)


class TransferConfirmIn(BaseModel):
    barcode: str


@api.post("/storage/transfers/confirm-putaway")
async def confirm_transfer_putaway(
    body: TransferConfirmIn,
    user: dict = Depends(require_role("admin", "manager", "operator")),
):
    """Scan a pallet barcode or destination location code to confirm physical putaway.
    Matches against all transfers with status='initiated'."""
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        scanned = body.barcode.strip()
        t = await conn.fetchrow(
            "SELECT * FROM transfers WHERE status='initiated' AND (pallet_code=$1 OR to_code=$1)",
            scanned,
        )
        if not t:
            raise HTTPException(404, f"No pending transfer found for '{scanned}'")
        ts = now_iso()
        await conn.execute(
            "UPDATE transfers SET status='confirmed', confirmed_at=$1, confirmed_by=$2 WHERE id=$3",
            ts, user["id"], t["id"],
        )
    return {
        "ok": True,
        "transfer_id": t["id"],
        "pallet_code": t["pallet_code"],
        "to_code": t["to_code"],
        "confirmed_at": ts,
    }


# ---------- Inbound ----------
@api.get("/inbound")
async def list_inbound(status: Optional[str] = None, user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        if status:
            orders = await conn.fetch(
                "SELECT i.*, v.name AS vendor_name FROM inbound i "
                "LEFT JOIN vendors v ON v.id = i.vendor_id WHERE i.status=$1 ORDER BY i.created_at DESC", status
            )
        else:
            orders = await conn.fetch(
                "SELECT i.*, v.name AS vendor_name FROM inbound i "
                "LEFT JOIN vendors v ON v.id = i.vendor_id ORDER BY i.created_at DESC"
            )
        if not orders:
            return []
        ids = [o["id"] for o in orders]
        items = await conn.fetch(
            "SELECT * FROM inbound_items WHERE inbound_id = ANY($1::text[])", ids
        )
    items_by_order: dict = {}
    for it in items:
        items_by_order.setdefault(it["inbound_id"], []).append(dict(it))
    result = []
    for o in orders:
        d = dict(o)
        d["items"] = items_by_order.get(d["id"], [])
        result.append(d)
    return result


@api.post("/inbound")
async def create_inbound(body: InboundIn, user: dict = Depends(require_role("admin", "manager"))):
    oid = str(uuid.uuid4())
    ts = now_iso()
    item_rows = []
    items_out = []
    for idx, item in enumerate(body.items):
        bag_color = item.bag_color.strip() if item.bag_color and item.bag_color.strip() else None
        if bag_color is not None and bag_color not in BAG_COLORS:
            raise HTTPException(status_code=400, detail=f"Bag color must be one of {sorted(BAG_COLORS)}")
        iid = str(uuid.uuid4())
        barcode = item.barcode or f"PLT-{body.po_number.replace(' ','-').upper()}-P{idx+1:02d}"
        item_rows.append((
            iid, oid, item.sku_id, item.qty, item.location_id, barcode,
            bag_color, item.batch_no, item.manufacture_date, item.expiry_date, False, None, None,
        ))
        items_out.append({
            "id": iid, "inbound_id": oid, "sku_id": item.sku_id, "qty": item.qty,
            "location_id": item.location_id, "barcode": barcode, "bag_color": bag_color,
            "batch_no": item.batch_no, "manufacture_date": item.manufacture_date,
            "expiry_date": item.expiry_date, "putaway_confirmed": False,
            "confirmed_at": None, "confirmed_by": None,
        })
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        async with conn.transaction():
            await conn.execute(
                "INSERT INTO inbound (id, po_number, supplier, vendor_id, expected_date, status, created_at, created_by) "
                "VALUES ($1,$2,$3,$4,$5,$6,$7,$8)",
                oid, body.po_number, body.supplier, body.vendor_id, body.expected_date, "pending", ts, user["name"],
            )
            await conn.executemany(
                "INSERT INTO inbound_items (id, inbound_id, sku_id, qty, location_id, barcode, "
                "bag_color, batch_no, manufacture_date, expiry_date, putaway_confirmed, confirmed_at, confirmed_by) "
                "VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10,$11,$12,$13)",
                item_rows,
            )
    return {
        "id": oid, "po_number": body.po_number, "supplier": body.supplier,
        "vendor_id": body.vendor_id, "vendor_name": None,
        "expected_date": body.expected_date, "status": "pending",
        "created_at": ts, "created_by": user["name"], "items": items_out,
    }


async def _execute_receive(conn, order: dict) -> None:
    received_at = now_iso()
    for item in order["items"]:
        # Actual bag color received on this line; fall back to the SKU default
        # so on-hand stock is always attributed to a color when one is known.
        color = item.get("bag_color")
        if not color:
            sku_def = await conn.fetchrow("SELECT bag_color FROM skus WHERE id=$1", item["sku_id"])
            color = sku_def["bag_color"] if sku_def else None
        # Match on color too, so pallets of different colors at the same
        # SKU/location stay as separate, correctly-coloured stock rows.
        existing = await conn.fetchrow(
            "SELECT id, qty FROM stock WHERE sku_id=$1 AND location_id=$2 "
            "AND COALESCE(bag_color,'')=COALESCE($3,'')",
            item["sku_id"], item["location_id"], color,
        )
        if existing:
            await conn.execute(
                "UPDATE stock SET qty=qty+$1, batch_no=$2, manufacture_date=$3, "
                "expiry_date=$4, received_date=$5, ref=$6 WHERE id=$7",
                item["qty"], item.get("batch_no"), item.get("manufacture_date"),
                item.get("expiry_date"), received_at, order["po_number"], existing["id"],
            )
        else:
            await conn.execute(
                "INSERT INTO stock (id, sku_id, location_id, qty, batch_no, manufacture_date, "
                "expiry_date, received_date, bag_color, ref, pallet_status, original_qty) "
                "VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10,$11,$12)",
                str(uuid.uuid4()), item["sku_id"], item["location_id"], item["qty"],
                item.get("batch_no"), item.get("manufacture_date"),
                item.get("expiry_date"), received_at, color, order["po_number"], "full", item["qty"],
            )
        await conn.execute(
            "UPDATE skus SET total_stock = total_stock + $1 WHERE id=$2", item["qty"], item["sku_id"]
        )
        await conn.execute(
            "UPDATE locations SET occupied = occupied + $1 WHERE id=$2", item["qty"], item["location_id"]
        )
        await conn.execute(
            "INSERT INTO movements (id, type, sku_id, location_id, qty, ref, batch_no, expiry_date, timestamp) "
            "VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9)",
            str(uuid.uuid4()), "in", item["sku_id"], item["location_id"], item["qty"],
            order["po_number"], item.get("batch_no"), item.get("expiry_date"), received_at,
        )
    await conn.execute(
        "UPDATE inbound SET status='completed', completed_at=$1 WHERE id=$2",
        received_at, order["id"],
    )


@api.post("/inbound/{order_id}/receive")
async def receive_inbound(order_id: str, user: dict = Depends(require_role("admin", "manager", "operator"))):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        async with conn.transaction():
            row = await conn.fetchrow("SELECT * FROM inbound WHERE id=$1", order_id)
            if not row:
                raise HTTPException(404, "Inbound order not found")
            if row["status"] == "completed":
                raise HTTPException(400, "Already completed")
            items = await conn.fetch("SELECT * FROM inbound_items WHERE inbound_id=$1", order_id)
            order = dict(row)
            order["items"] = [dict(it) for it in items]
            await _execute_receive(conn, order)
    return {"ok": True}


@api.post("/inbound/{order_id}/putaway-scan")
async def putaway_scan(order_id: str, body: PutawayScanIn, user: dict = Depends(require_role("admin", "manager", "operator"))):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        async with conn.transaction():
            row = await conn.fetchrow("SELECT * FROM inbound WHERE id=$1", order_id)
            if not row:
                raise HTTPException(404, "Inbound order not found")
            if row["status"] == "completed":
                raise HTTPException(400, "Order already completed")
            items = await conn.fetch("SELECT * FROM inbound_items WHERE inbound_id=$1", order_id)
            matched = next((it for it in items if it["barcode"] == body.barcode.strip()), None)
            if not matched:
                raise HTTPException(404, f"Barcode '{body.barcode}' not found in this order")
            if matched["putaway_confirmed"]:
                raise HTTPException(400, "Pallet already confirmed")
            confirmed_at = now_iso()
            await conn.execute(
                "UPDATE inbound_items SET putaway_confirmed=TRUE, confirmed_at=$1, confirmed_by=$2 WHERE id=$3",
                confirmed_at, user["name"], matched["id"],
            )
            all_confirmed = all(
                it["putaway_confirmed"] or it["id"] == matched["id"] for it in items
            )
            if all_confirmed:
                fresh_items = await conn.fetch("SELECT * FROM inbound_items WHERE inbound_id=$1", order_id)
                order = dict(row)
                order["items"] = [dict(it) for it in fresh_items]
                await _execute_receive(conn, order)
    return {
        "ok": True, "barcode": body.barcode,
        "item_index": next(i for i, it in enumerate(items) if it["id"] == matched["id"]),
        "all_confirmed": all_confirmed,
    }


# ---------- GRN / Pick List ----------
@api.get("/inbound/{order_id}/grn")
async def inbound_grn(order_id: str, user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        row = await conn.fetchrow("SELECT * FROM inbound WHERE id=$1", order_id)
        if not row:
            raise HTTPException(404, "Not found")
        items = await conn.fetch("SELECT * FROM inbound_items WHERE inbound_id=$1", order_id)
        enriched = []
        for it in items:
            sku = await conn.fetchrow(
                "SELECT sku_code, name, unit, bag_color, weight_per_bag, bags_per_pallet, dimensions "
                "FROM skus WHERE id=$1",
                it["sku_id"],
            )
            loc = await conn.fetchrow("SELECT code FROM locations WHERE id=$1", it["location_id"])
            enriched.append({**dict(it), "sku": dict(sku) if sku else None, "location": dict(loc) if loc else None})
    order = dict(row)
    order["items"] = enriched
    return order


@api.get("/outbound/{order_id}/picklist")
async def outbound_picklist(order_id: str, user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        row = await conn.fetchrow("SELECT * FROM outbound WHERE id=$1", order_id)
        if not row:
            raise HTTPException(404, "Not found")
        items = await conn.fetch("SELECT * FROM outbound_items WHERE outbound_id=$1", order_id)
        enriched = []
        for it in items:
            sku = await conn.fetchrow(
                "SELECT sku_code, name, unit, bag_color, weight_per_bag, bags_per_pallet, dimensions "
                "FROM skus WHERE id=$1",
                it["sku_id"],
            )
            pallets = await conn.fetch(
                "SELECT * FROM stock WHERE sku_id=$1 AND qty > 0 ORDER BY expiry_date NULLS LAST",
                it["sku_id"],
            )
            pick_locs = []
            remaining = it["qty"]
            for p in pallets:
                if remaining <= 0:
                    break
                loc = await conn.fetchrow("SELECT code FROM locations WHERE id=$1", p["location_id"])
                take = min(p["qty"], remaining)
                pick_locs.append({"location": dict(loc) if loc else None, "qty": take,
                                  "batch_no": p["batch_no"], "expiry_date": p["expiry_date"]})
                remaining -= take
            enriched.append({**dict(it), "sku": dict(sku) if sku else None, "pick_locations": pick_locs})
    order = dict(row)
    order["items"] = enriched
    return order


# ---------- Outbound ----------
@api.get("/outbound")
async def list_outbound(status: Optional[str] = None, user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        if status:
            orders = await conn.fetch(
                "SELECT o.*, c.name AS customer_name FROM outbound o "
                "LEFT JOIN customers c ON c.id = o.customer_id WHERE o.status=$1 ORDER BY o.created_at DESC", status
            )
        else:
            orders = await conn.fetch(
                "SELECT o.*, c.name AS customer_name FROM outbound o "
                "LEFT JOIN customers c ON c.id = o.customer_id ORDER BY o.created_at DESC"
            )
        if not orders:
            return []
        ids = [o["id"] for o in orders]
        items = await conn.fetch("SELECT * FROM outbound_items WHERE outbound_id = ANY($1::text[])", ids)
        picks = await conn.fetch(
            "SELECT outbound_id, sku_id, COALESCE(SUM(qty),0) AS picked FROM outbound_picks "
            "WHERE outbound_id = ANY($1::text[]) GROUP BY outbound_id, sku_id", ids
        )
    remaining_by_order_sku: dict = {}
    for p in picks:
        remaining_by_order_sku[(p["outbound_id"], p["sku_id"])] = p["picked"]
    items_by_order: dict = {}
    # Allocate the SKU-level picked total across that SKU's order lines in turn,
    # so a SKU spread over multiple lines reports accurate per-line progress.
    for it in sorted(items, key=lambda r: r["id"]):
        d = dict(it)
        key = (it["outbound_id"], it["sku_id"])
        avail = remaining_by_order_sku.get(key, 0)
        alloc = min(avail, it["qty"])
        d["picked_qty"] = alloc
        remaining_by_order_sku[key] = avail - alloc
        items_by_order.setdefault(it["outbound_id"], []).append(d)
    result = []
    for o in orders:
        d = dict(o)
        d["items"] = items_by_order.get(d["id"], [])
        result.append(d)
    return result


@api.post("/outbound")
async def create_outbound(body: OutboundIn, user: dict = Depends(require_role("admin", "manager"))):
    oid = str(uuid.uuid4())
    ts = now_iso()
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        await conn.execute(
            "INSERT INTO outbound (id, so_number, customer, customer_id, status, created_at, created_by) VALUES ($1,$2,$3,$4,$5,$6,$7)",
            oid, body.so_number, body.customer, body.customer_id, "pending", ts, user["name"],
        )
        item_rows = [(str(uuid.uuid4()), oid, it.sku_id, it.qty) for it in body.items]
        await conn.executemany(
            "INSERT INTO outbound_items (id, outbound_id, sku_id, qty) VALUES ($1,$2,$3,$4)", item_rows
        )
    items_out = [{"id": r[0], "outbound_id": oid, "sku_id": r[2], "qty": r[3]} for r in item_rows]
    return {
        "id": oid, "so_number": body.so_number, "customer": body.customer,
        "customer_id": body.customer_id, "customer_name": None,
        "status": "pending", "created_at": ts, "created_by": user["name"], "items": items_out,
    }


@api.post("/outbound/{order_id}/advance")
async def advance_outbound(order_id: str, user: dict = Depends(require_role("admin", "manager", "operator"))):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        async with conn.transaction():
            row = await conn.fetchrow("SELECT * FROM outbound WHERE id=$1", order_id)
            if not row:
                raise HTTPException(404, "Outbound order not found")
            flow = {"pending": "picking", "picking": "packing", "packing": "shipped"}
            nxt = flow.get(row["status"])
            if not nxt:
                raise HTTPException(400, "Order is already shipped")
            if nxt == "shipped":
                items = await conn.fetch("SELECT * FROM outbound_items WHERE outbound_id=$1", order_id)
                for item in items:
                    remaining = item["qty"]
                    # Partial pallets (oldest first) → full pallets (FIFO by received_date)
                    stocks = await conn.fetch(
                        """SELECT * FROM stock
                           WHERE sku_id=$1 AND qty > 0
                             AND COALESCE(pallet_status,'full') NOT IN ('damaged','quality_hold','empty')
                           ORDER BY
                             CASE WHEN COALESCE(pallet_status,'full')='partial' THEN 0 ELSE 1 END,
                             COALESCE(partial_since, received_date, ''),
                             received_date NULLS LAST""",
                        item["sku_id"]
                    )
                    for s in stocks:
                        if remaining <= 0:
                            break
                        take = min(s["qty"], remaining)
                        await conn.execute(
                            "UPDATE stock SET qty=qty-$1 WHERE id=$2", take, s["id"]
                        )
                        await conn.execute(
                            "UPDATE locations SET occupied=GREATEST(0, occupied-$1) WHERE id=$2",
                            take, s["location_id"],
                        )
                        await conn.execute(
                            "INSERT INTO movements (id, type, sku_id, location_id, qty, ref, timestamp) "
                            "VALUES ($1,$2,$3,$4,$5,$6,$7)",
                            str(uuid.uuid4()), "out", item["sku_id"], s["location_id"],
                            take, row["so_number"], now_iso(),
                        )
                        remaining -= take
                        # Update pallet status after deduction
                        new_qty = s["qty"] - take
                        ts = now_iso()
                        if new_qty <= 0:
                            await conn.execute(
                                "UPDATE stock SET pallet_status='empty' WHERE id=$1", s["id"]
                            )
                        else:
                            orig = s["original_qty"] if s["original_qty"] is not None else (s["qty"] + take)
                            if new_qty < orig:
                                await conn.execute(
                                    "UPDATE stock SET pallet_status='partial', "
                                    "partial_since=COALESCE(partial_since,$1) WHERE id=$2",
                                    ts, s["id"]
                                )
                                await conn.execute(
                                    "INSERT INTO partial_pallet_history "
                                    "(id,stock_id,transaction_type,outbound_id,so_number,"
                                    "user_name,qty_picked,remaining_qty,prev_location_id,timestamp) "
                                    "VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10)",
                                    str(uuid.uuid4()), s["id"], "outbound_ship",
                                    order_id, row["so_number"],
                                    user.get("name", ""), take, new_qty, s["location_id"], ts
                                )
                    deducted = item["qty"] - remaining
                    await conn.execute(
                        "UPDATE skus SET total_stock=GREATEST(0, total_stock-$1) WHERE id=$2",
                        deducted, item["sku_id"],
                    )
            await conn.execute("UPDATE outbound SET status=$1 WHERE id=$2", nxt, order_id)
    return {"ok": True, "status": nxt}


# ── Partial Pallet Management ──────────────────────────────────────────────────

class PartialPalletStatusIn(BaseModel):
    status: str  # partial, damaged, quality_hold


@api.get("/partial-pallets/dashboard")
async def partial_pallets_dashboard(
    user: dict = Depends(require_role("admin", "manager", "operator")),
):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        stats = await conn.fetchrow("""
            SELECT
                COUNT(*) FILTER (WHERE COALESCE(pallet_status,'full')='partial') AS total_partial,
                COALESCE(SUM(qty) FILTER (WHERE COALESCE(pallet_status,'full')='partial'), 0) AS total_qty,
                COUNT(DISTINCT sku_id) FILTER (WHERE COALESCE(pallet_status,'full')='partial') AS product_count,
                MIN(COALESCE(partial_since, received_date))
                    FILTER (WHERE COALESCE(pallet_status,'full')='partial') AS oldest,
                COUNT(*) FILTER (
                    WHERE COALESCE(pallet_status,'full')='partial'
                      AND expiry_date IS NOT NULL
                      AND expiry_date <= (CURRENT_DATE + INTERVAL '30 days')::TEXT
                ) AS near_expiry_count,
                AVG(
                    CASE WHEN original_qty > 0 THEN qty::float / original_qty ELSE NULL END
                ) FILTER (WHERE COALESCE(pallet_status,'full')='partial') AS avg_utilization
            FROM stock
            WHERE qty > 0
        """)
        oldest_days = None
        if stats["oldest"]:
            try:
                oldest_dt = datetime.fromisoformat(stats["oldest"].replace("Z", "+00:00"))
                oldest_days = (datetime.now(timezone.utc) - oldest_dt).days
            except Exception:
                pass
        return {
            "total_partial": stats["total_partial"] or 0,
            "total_qty": int(stats["total_qty"] or 0),
            "product_count": stats["product_count"] or 0,
            "oldest_days": oldest_days,
            "near_expiry_count": stats["near_expiry_count"] or 0,
            "utilization_pct": round((stats["avg_utilization"] or 0) * 100, 1),
        }


@api.get("/partial-pallets")
async def list_partial_pallets(
    sku_id: Optional[str] = None,
    batch_no: Optional[str] = None,
    status: Optional[str] = None,
    expiry_before: Optional[str] = None,
    location_id: Optional[str] = None,
    user: dict = Depends(require_role("admin", "manager", "operator")),
):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        q = """
            SELECT s.id, s.sku_id, s.location_id, s.qty, s.original_qty,
                   s.batch_no, s.manufacture_date, s.expiry_date, s.received_date,
                   s.pallet_code, s.pallet_status, s.partial_since,
                   sk.sku_code, sk.name AS sku_name, sk.unit,
                   l.code AS location_code, l.zone, l.zone_name,
                   l.row_label, l.lane_number, l.level
            FROM stock s
            JOIN skus sk ON sk.id = s.sku_id
            LEFT JOIN locations l ON l.id = s.location_id
            WHERE COALESCE(s.pallet_status, 'full') IN ('partial','damaged','quality_hold')
              AND s.qty > 0
        """
        params: list = []
        if sku_id:
            params.append(sku_id)
            q += f" AND s.sku_id=${len(params)}"
        if batch_no:
            params.append(f"%{batch_no}%")
            q += f" AND s.batch_no ILIKE ${len(params)}"
        if status:
            params.append(status)
            q += f" AND s.pallet_status=${len(params)}"
        if expiry_before:
            params.append(expiry_before)
            q += f" AND s.expiry_date <= ${len(params)}"
        if location_id:
            params.append(location_id)
            q += f" AND s.location_id=${len(params)}"
        q += " ORDER BY COALESCE(s.partial_since, s.received_date) NULLS LAST"
        rows = await conn.fetch(q, *params)
        now_dt = datetime.now(timezone.utc)
        result = []
        for r in rows:
            d = dict(r)
            if d.get("partial_since"):
                try:
                    ps = datetime.fromisoformat(d["partial_since"].replace("Z", "+00:00"))
                    d["partial_age_days"] = (now_dt - ps).days
                except Exception:
                    d["partial_age_days"] = None
            else:
                d["partial_age_days"] = None
            result.append(d)
        return result


@api.get("/partial-pallets/{stock_id}/history")
async def get_partial_pallet_history(
    stock_id: str,
    user: dict = Depends(require_role("admin", "manager", "operator")),
):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        rows = await conn.fetch(
            "SELECT * FROM partial_pallet_history WHERE stock_id=$1 ORDER BY timestamp DESC",
            stock_id,
        )
        return [dict(r) for r in rows]


@api.patch("/partial-pallets/{stock_id}/status")
async def update_partial_pallet_status(
    stock_id: str,
    body: PartialPalletStatusIn,
    user: dict = Depends(require_role("admin", "manager")),
):
    valid = {"partial", "damaged", "quality_hold"}
    if body.status not in valid:
        raise HTTPException(400, f"status must be one of: {', '.join(sorted(valid))}")
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        row = await conn.fetchrow("SELECT id FROM stock WHERE id=$1", stock_id)
        if not row:
            raise HTTPException(404, "Pallet not found")
        await conn.execute(
            "UPDATE stock SET pallet_status=$1 WHERE id=$2", body.status, stock_id
        )
    return {"ok": True}


async def _resolve_pallet(conn, barcode: str):
    """Resolve a scanned pallet barcode to (sku_id, qty). Checks putaway stock
    pallet codes first, then the inbound item barcode generated at receipt.
    Only pallets that are actually in the warehouse are pickable: live stock
    (qty > 0) or an inbound line whose order has been received (status
    'completed')."""
    s = await conn.fetchrow(
        "SELECT sku_id, qty FROM stock WHERE pallet_code=$1 AND qty > 0 ORDER BY qty DESC", barcode
    )
    if s:
        return s["sku_id"], s["qty"]
    ii = await conn.fetchrow(
        "SELECT ii.sku_id, ii.qty FROM inbound_items ii "
        "JOIN inbound i ON i.id = ii.inbound_id "
        "WHERE ii.barcode=$1 AND i.status='completed' ORDER BY ii.id LIMIT 1", barcode
    )
    if ii:
        return ii["sku_id"], ii["qty"]
    return None


@api.post("/outbound/{order_id}/pick-scan")
async def outbound_pick_scan(
    order_id: str, body: PickScanIn,
    user: dict = Depends(require_role("admin", "manager", "operator")),
):
    barcode = (body.barcode or "").strip()
    if not barcode:
        raise HTTPException(400, "Barcode is required")
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        async with conn.transaction():
            order = await conn.fetchrow("SELECT * FROM outbound WHERE id=$1 FOR UPDATE", order_id)
            if not order:
                raise HTTPException(404, "Outbound order not found")
            if order["status"] not in ("pending", "picking"):
                raise HTTPException(400, "Order is already past picking")

            resolved = await _resolve_pallet(conn, barcode)
            if not resolved:
                raise HTTPException(404, f"Unknown pallet barcode: {barcode}")
            sku_id, pallet_qty = resolved

            items = await conn.fetch("SELECT * FROM outbound_items WHERE outbound_id=$1", order_id)
            ordered_by_sku: dict = {}
            for it in items:
                ordered_by_sku[it["sku_id"]] = ordered_by_sku.get(it["sku_id"], 0) + it["qty"]
            if sku_id not in ordered_by_sku:
                sku = await conn.fetchrow("SELECT sku_code FROM skus WHERE id=$1", sku_id)
                code = sku["sku_code"] if sku else sku_id
                raise HTTPException(400, f"Scanned pallet ({code}) is not part of this order")

            dup = await conn.fetchrow(
                "SELECT id FROM outbound_picks WHERE outbound_id=$1 AND barcode=$2", order_id, barcode
            )
            if dup:
                raise HTTPException(400, "This pallet has already been scanned for this order")

            try:
                await conn.execute(
                    "INSERT INTO outbound_picks (id, outbound_id, barcode, sku_id, qty, scanned_at, scanned_by) "
                    "VALUES ($1,$2,$3,$4,$5,$6,$7)",
                    str(uuid.uuid4()), order_id, barcode, sku_id, pallet_qty, now_iso(), user["name"],
                )
            except asyncpg.UniqueViolationError:
                raise HTTPException(400, "This pallet has already been scanned for this order")

            pick_rows = await conn.fetch(
                "SELECT sku_id, COALESCE(SUM(qty),0) AS picked FROM outbound_picks "
                "WHERE outbound_id=$1 GROUP BY sku_id", order_id
            )
            picked_by_sku = {r["sku_id"]: r["picked"] for r in pick_rows}
            fully_picked = all(
                picked_by_sku.get(sid, 0) >= qty for sid, qty in ordered_by_sku.items()
            )

            new_status = "packing" if fully_picked else "picking"
            if order["status"] != new_status:
                await conn.execute("UPDATE outbound SET status=$1 WHERE id=$2", new_status, order_id)

            sku = await conn.fetchrow("SELECT sku_code FROM skus WHERE id=$1", sku_id)
    ordered_qty = ordered_by_sku[sku_id]
    return {
        "ok": True,
        "status": new_status,
        "fully_picked": fully_picked,
        "sku_code": sku["sku_code"] if sku else sku_id,
        "picked_qty": min(picked_by_sku.get(sku_id, 0), ordered_qty),
        "ordered_qty": ordered_qty,
    }


# ---------- Dashboard ----------
@api.get("/dashboard/summary")
async def dashboard_summary(user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        total_skus = await conn.fetchval("SELECT COUNT(*) FROM skus")
        total_locations = await conn.fetchval("SELECT COUNT(*) FROM locations")
        sku_rows = await conn.fetch("SELECT total_stock, unit_price, reorder_level FROM skus")
        total_value = sum(s["total_stock"] * s["unit_price"] for s in sku_rows)
        low_stock = sum(1 for s in sku_rows if 0 < s["total_stock"] <= s["reorder_level"])
        out_of_stock = sum(1 for s in sku_rows if s["total_stock"] <= 0)
        today = datetime.now(timezone.utc).date().isoformat()
        inbound_today = await conn.fetchval(
            "SELECT COUNT(*) FROM movements WHERE type='in' AND timestamp >= $1", today
        )
        outbound_today = await conn.fetchval(
            "SELECT COUNT(*) FROM movements WHERE type='out' AND timestamp >= $1", today
        )
        pending_inbound = await conn.fetchval("SELECT COUNT(*) FROM inbound WHERE status='pending'")
        pending_outbound = await conn.fetchval(
            "SELECT COUNT(*) FROM outbound WHERE status != 'shipped'"
        )
        cap_row = await conn.fetchrow(
            "SELECT SUM(capacity) AS cap, SUM(occupied) AS occ FROM locations"
        )
        utilization = 0
        if cap_row and cap_row["cap"]:
            utilization = round(cap_row["occ"] / cap_row["cap"] * 100, 1)
        recent = await conn.fetch(
            "SELECT m.*, sk.sku_code, sk.name AS sk_name, l.code AS loc_code "
            "FROM movements m "
            "LEFT JOIN skus sk ON sk.id = m.sku_id "
            "LEFT JOIN locations l ON l.id = m.location_id "
            "ORDER BY m.timestamp DESC LIMIT 8"
        )
    recent_out = []
    for m in recent:
        d = dict(m)
        d["sku"] = {"sku_code": d.pop("sku_code", None), "name": d.pop("sk_name", None)}
        d["location"] = {"code": d.pop("loc_code", None)}
        recent_out.append(d)
    return {
        "total_skus": total_skus, "total_locations": total_locations,
        "stock_value": round(total_value, 2), "low_stock": low_stock,
        "out_of_stock": out_of_stock, "inbound_today": inbound_today,
        "outbound_today": outbound_today, "pending_inbound": pending_inbound,
        "pending_outbound": pending_outbound, "utilization": utilization,
        "recent_movements": recent_out,
    }


@api.get("/dashboard/movements")
async def movement_trends(days: int = 14, user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    start = (datetime.now(timezone.utc) - timedelta(days=days)).date().isoformat()
    async with pool.acquire() as conn:
        rows = await conn.fetch(
            "SELECT type, qty, timestamp FROM movements WHERE timestamp >= $1", start
        )
    buckets: dict = {}
    for m in rows:
        day = m["timestamp"][:10]
        b = buckets.setdefault(day, {"date": day, "inbound": 0, "outbound": 0})
        if m["type"] == "in":
            b["inbound"] += m["qty"]
        else:
            b["outbound"] += m["qty"]
    out = []
    for i in range(days, -1, -1):
        d = (datetime.now(timezone.utc) - timedelta(days=i)).date().isoformat()
        out.append(buckets.get(d, {"date": d, "inbound": 0, "outbound": 0}))
    return out


# ---------- Reports ----------
@api.get("/reports/top-skus")
async def top_skus(limit: int = 10, user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        rows = await conn.fetch(
            "SELECT m.sku_id, SUM(m.qty) AS moved, sk.sku_code, sk.name, sk.category "
            "FROM movements m JOIN skus sk ON sk.id = m.sku_id "
            "GROUP BY m.sku_id, sk.sku_code, sk.name, sk.category "
            "ORDER BY moved DESC LIMIT $1", limit,
        )
    return [{"sku": {"sku_code": r["sku_code"], "name": r["name"], "category": r["category"]}, "moved": r["moved"]} for r in rows]


@api.get("/reports/category-distribution")
async def category_dist(user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        rows = await conn.fetch(
            "SELECT category, COUNT(*) AS count, SUM(total_stock) AS stock FROM skus "
            "GROUP BY category ORDER BY count DESC"
        )
    return [{"category": r["category"], "count": r["count"], "stock": r["stock"] or 0} for r in rows]


@api.post("/reports/ai-insights")
async def ai_insights(user: dict = Depends(get_user)):
    summary = await dashboard_summary(user)
    top = await top_skus(5, user)
    cats = await category_dist(user)
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        low_stock_skus = await conn.fetch(
            "SELECT sku_code, name, total_stock, reorder_level FROM skus "
            "WHERE total_stock <= reorder_level ORDER BY total_stock LIMIT 10"
        )
        weight_rows = await conn.fetch(
            "SELECT category, COALESCE(SUM(weight_per_bag * total_stock), 0) AS weight FROM skus "
            "GROUP BY category ORDER BY weight DESC"
        )
    total_weight = sum(float(w["weight"] or 0) for w in weight_rows)
    weight_lines = chr(10).join(
        [f"  - {w['category']}: {float(w['weight'] or 0):,.1f} kg" for w in weight_rows]
    )
    context = f"""
WAREHOUSE PERFORMANCE SNAPSHOT
- Total SKUs: {summary['total_skus']}
- Stock Value: ${summary['stock_value']}
- Low Stock SKUs: {summary['low_stock']}
- Out of Stock: {summary['out_of_stock']}
- Storage Utilization: {summary['utilization']}%
- Pending Inbound: {summary['pending_inbound']}
- Pending Outbound: {summary['pending_outbound']}

LOW STOCK ITEMS:
{chr(10).join([f"  - {s['sku_code']} {s['name']}: {s['total_stock']} on hand (reorder at {s['reorder_level']})" for s in low_stock_skus])}

TOP MOVING SKUs:
{chr(10).join([f"  - {t['sku']['sku_code']} {t['sku']['name']}: {t['moved']} units moved" for t in top])}

CATEGORY DISTRIBUTION:
{chr(10).join([f"  - {c['category']}: {c['count']} SKUs, {c['stock']} units" for c in cats])}

STORED WEIGHT (weight_per_bag x on-hand bags):
- Total stored weight: {total_weight:,.1f} kg
{weight_lines}
"""
    try:
        _client, _model = _get_anthropic_client()
        _resp = _client.messages.create(
            model=_model,
            max_tokens=1024,
            system=(
                "You are a senior cold-chain warehouse operations analyst overseeing a -20°C drive-in racking facility (LIFO). "
                "Analyze the snapshot and produce a sharp, tactical report with: 1) a 2-line executive summary, "
                "2) 3 prioritized actionable insights (with bullet markers '>'), 3) a brief risk assessment. "
                "Use crisp, technical language. Format output as plain text with section headers in CAPS. Keep total under 220 words."
            ),
            messages=[{"role": "user", "content": context}],
        )
        response = _resp.content[0].text
        return {"insights": response, "generated_at": now_iso()}
    except Exception as e:
        logging.exception("AI insights failed")
        err = str(e)
        if "not configured" in err.lower() or "not_found_error" in err.lower():
            msg = (
                "AI assistant requires a valid Anthropic API key.\n\n"
                "To enable: go to your Anthropic Console (console.anthropic.com), "
                "generate an API key, then update the ANTHROPIC_API_KEY secret in this Replit project "
                "with that key (it must be a direct Anthropic key, not a Replit-managed one)."
            )
        else:
            msg = f"AI service unavailable: {err}"
        return {"insights": msg, "generated_at": now_iso()}


# ─── Extended Reports ─────────────────────────────────────────────────────────

@api.get("/reports/stock-on-hand")
async def report_stock_on_hand(view: str = "sku", zone: Optional[str] = None, user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        if view == "location":
            q = ("SELECT l.code AS location, l.zone, COALESCE(l.zone_name,'') AS zone_name, "
                 "l.row_label, l.lane_number, l.level, "
                 "sk.sku_code, sk.name AS sku_name, s.batch_no, s.expiry_date, s.bag_color, SUM(s.qty) AS qty "
                 "FROM stock s JOIN skus sk ON sk.id=s.sku_id JOIN locations l ON l.id=s.location_id "
                 "WHERE s.qty>0 ")
            args: list = []
            if zone:
                args.append(zone)
                q += f"AND l.zone=${len(args)} "
            q += ("GROUP BY l.code,l.zone,l.zone_name,l.row_label,l.lane_number,l.level,"
                  "sk.sku_code,sk.name,s.batch_no,s.expiry_date,s.bag_color ORDER BY l.zone,l.code")
            rows = await conn.fetch(q, *args)
        elif view == "batch":
            rows = await conn.fetch(
                "SELECT sk.sku_code, sk.name AS sku_name, s.batch_no, s.expiry_date, s.manufacture_date, "
                "s.bag_color, SUM(s.qty) AS qty, COUNT(DISTINCT s.location_id) AS location_count "
                "FROM stock s JOIN skus sk ON sk.id=s.sku_id "
                "WHERE s.qty>0 AND s.batch_no IS NOT NULL "
                "GROUP BY sk.sku_code,sk.name,s.batch_no,s.expiry_date,s.manufacture_date,s.bag_color "
                "ORDER BY sk.sku_code,s.expiry_date"
            )
        else:
            rows = await conn.fetch(
                "SELECT sk.sku_code, sk.name AS sku_name, sk.category, sk.unit, "
                "SUM(s.qty) AS total_qty, COUNT(DISTINCT s.location_id) AS locations, "
                "MIN(s.expiry_date) AS earliest_expiry "
                "FROM stock s JOIN skus sk ON sk.id=s.sku_id "
                "WHERE s.qty>0 "
                "GROUP BY sk.sku_code,sk.name,sk.category,sk.unit ORDER BY sk.sku_code"
            )
    return [dict(r) for r in rows]


@api.get("/reports/expiry")
async def report_expiry(days: int = 30, user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    today = date.today().isoformat()
    async with pool.acquire() as conn:
        if days == 0:
            rows = await conn.fetch(
                "SELECT sk.sku_code, sk.name AS sku_name, s.batch_no, s.expiry_date, "
                "l.code AS location, SUM(s.qty) AS qty "
                "FROM stock s JOIN skus sk ON sk.id=s.sku_id JOIN locations l ON l.id=s.location_id "
                "WHERE s.qty>0 AND s.expiry_date IS NOT NULL AND s.expiry_date<$1 "
                "GROUP BY sk.sku_code,sk.name,s.batch_no,s.expiry_date,l.code "
                "ORDER BY s.expiry_date", today
            )
        else:
            future = (date.today() + timedelta(days=days)).isoformat()
            rows = await conn.fetch(
                "SELECT sk.sku_code, sk.name AS sku_name, s.batch_no, s.expiry_date, "
                "l.code AS location, SUM(s.qty) AS qty "
                "FROM stock s JOIN skus sk ON sk.id=s.sku_id JOIN locations l ON l.id=s.location_id "
                "WHERE s.qty>0 AND s.expiry_date IS NOT NULL AND s.expiry_date>=$1 AND s.expiry_date<=$2 "
                "GROUP BY sk.sku_code,sk.name,s.batch_no,s.expiry_date,l.code "
                "ORDER BY s.expiry_date", today, future
            )
    return [dict(r) for r in rows]


@api.get("/reports/grn")
async def report_grn(
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    vendor_id: Optional[str] = None,
    user: dict = Depends(get_user),
):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        conds = ["i.status='completed'"]
        args: list = []
        if date_from:
            args.append(date_from)
            conds.append(f"i.created_at>=${len(args)}")
        if date_to:
            args.append(date_to + "T23:59:59")
            conds.append(f"i.created_at<=${len(args)}")
        if vendor_id:
            args.append(vendor_id)
            conds.append(f"i.vendor_id=${len(args)}")
        where = " AND ".join(conds)
        rows = await conn.fetch(
            f"SELECT i.id, i.po_number, COALESCE(v.name,i.supplier) AS vendor_name, "
            f"i.expected_date, i.completed_at, i.created_at, i.created_by, "
            f"COUNT(ii.id) AS line_count, COALESCE(SUM(ii.qty),0) AS total_qty "
            f"FROM inbound i "
            f"LEFT JOIN vendors v ON v.id=i.vendor_id "
            f"LEFT JOIN inbound_items ii ON ii.inbound_id=i.id "
            f"WHERE {where} "
            f"GROUP BY i.id,i.po_number,v.name,i.supplier,i.expected_date,i.completed_at,i.created_at,i.created_by "
            f"ORDER BY i.created_at DESC", *args
        )
    return [dict(r) for r in rows]


@api.get("/reports/pending-receipts")
async def report_pending_receipts(user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    today = date.today().isoformat()
    async with pool.acquire() as conn:
        rows = await conn.fetch(
            "SELECT i.id, i.po_number, COALESCE(v.name,i.supplier) AS vendor_name, "
            "i.expected_date, i.created_at, i.created_by, "
            "COALESCE(SUM(ii.qty),0) AS expected_qty, COUNT(ii.id) AS line_count, "
            "CASE WHEN i.expected_date<$1 THEN 'overdue' ELSE 'on-time' END AS timeliness "
            "FROM inbound i "
            "LEFT JOIN vendors v ON v.id=i.vendor_id "
            "LEFT JOIN inbound_items ii ON ii.inbound_id=i.id "
            "WHERE i.status='pending' "
            "GROUP BY i.id,i.po_number,v.name,i.supplier,i.expected_date,i.created_at,i.created_by "
            "ORDER BY i.expected_date", today
        )
    return [dict(r) for r in rows]


@api.get("/reports/shipments")
async def report_shipments(
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    customer_id: Optional[str] = None,
    status: Optional[str] = None,
    user: dict = Depends(get_user),
):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        conds: list = ["1=1"]
        args: list = []
        if date_from:
            args.append(date_from)
            conds.append(f"o.created_at>=${len(args)}")
        if date_to:
            args.append(date_to + "T23:59:59")
            conds.append(f"o.created_at<=${len(args)}")
        if customer_id:
            args.append(customer_id)
            conds.append(f"o.customer_id=${len(args)}")
        if status:
            args.append(status)
            conds.append(f"o.status=${len(args)}")
        where = " AND ".join(conds)
        rows = await conn.fetch(
            f"SELECT o.id, o.so_number, COALESCE(c.name,o.customer) AS customer_name, "
            f"o.status, o.created_at, o.created_by, "
            f"COUNT(oi.id) AS line_count, COALESCE(SUM(oi.qty),0) AS total_qty "
            f"FROM outbound o "
            f"LEFT JOIN customers c ON c.id=o.customer_id "
            f"LEFT JOIN outbound_items oi ON oi.outbound_id=o.id "
            f"WHERE {where} "
            f"GROUP BY o.id,o.so_number,c.name,o.customer,o.status,o.created_at,o.created_by "
            f"ORDER BY o.created_at DESC", *args
        )
    return [dict(r) for r in rows]


@api.get("/reports/pick-performance")
async def report_pick_performance(user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        rows = await conn.fetch(
            "SELECT o.so_number, COALESCE(c.name,o.customer) AS customer_name, "
            "o.status, o.created_at, sk.sku_code, sk.name AS sku_name, "
            "SUM(oi.qty) AS ordered_qty, COALESCE(op.picked_qty,0) AS picked_qty, "
            "ROUND(COALESCE(op.picked_qty,0)*100.0/NULLIF(SUM(oi.qty),0),1) AS accuracy_pct "
            "FROM outbound_items oi "
            "JOIN outbound o ON o.id=oi.outbound_id "
            "JOIN skus sk ON sk.id=oi.sku_id "
            "LEFT JOIN customers c ON c.id=o.customer_id "
            "LEFT JOIN ("
            "  SELECT outbound_id,sku_id,SUM(qty) AS picked_qty "
            "  FROM outbound_picks GROUP BY outbound_id,sku_id"
            ") op ON op.outbound_id=oi.outbound_id AND op.sku_id=oi.sku_id "
            "GROUP BY o.id,o.so_number,c.name,o.customer,o.status,o.created_at,"
            "sk.sku_code,sk.name,op.picked_qty "
            "ORDER BY o.created_at DESC"
        )
    return [dict(r) for r in rows]


@api.get("/reports/stock-movement")
async def report_stock_movement(
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    mov_type: Optional[str] = None,
    sku_search: Optional[str] = None,
    user: dict = Depends(get_user),
):
    pool = await _db.get_pool()
    sku_like = f"%{sku_search.strip()}%" if sku_search and sku_search.strip() else None
    async with pool.acquire() as conn:
        conds: list = ["1=1"]
        args: list = []
        if date_from:
            args.append(date_from)
            conds.append(f"m.timestamp>=${len(args)}")
        if date_to:
            args.append(date_to + "T23:59:59")
            conds.append(f"m.timestamp<=${len(args)}")
        if mov_type and mov_type != "transfer":
            args.append(mov_type)
            conds.append(f"m.type=${len(args)}")
        if sku_like:
            args.append(sku_like)
            conds.append(f"(sk.sku_code ILIKE ${len(args)} OR sk.name ILIKE ${len(args)})")
        where = " AND ".join(conds)
        mov_rows = [] if mov_type == "transfer" else await conn.fetch(
            f"SELECT m.type, m.timestamp, sk.sku_code, sk.name AS sku_name, "
            f"l.code AS location, m.qty, m.ref, m.batch_no "
            f"FROM movements m JOIN skus sk ON sk.id=m.sku_id JOIN locations l ON l.id=m.location_id "
            f"WHERE {where} ORDER BY m.timestamp DESC LIMIT 1000", *args
        )
        t_conds: list = ["status='confirmed'"]
        t_args: list = []
        if date_from:
            t_args.append(date_from)
            t_conds.append(f"transferred_at>=${len(t_args)}")
        if date_to:
            t_args.append(date_to + "T23:59:59")
            t_conds.append(f"transferred_at<=${len(t_args)}")
        if sku_like:
            t_args.append(sku_like)
            t_conds.append(f"(sku_code ILIKE ${len(t_args)} OR sku_name ILIKE ${len(t_args)})")
        t_where = " AND ".join(t_conds)
        xfer_rows = [] if (mov_type and mov_type != "transfer") else await conn.fetch(
            f"SELECT 'transfer' AS type, transferred_at AS timestamp, sku_code, sku_name, "
            f"from_code||' → '||to_code AS location, qty, id AS ref, NULL::text AS batch_no "
            f"FROM transfers WHERE {t_where} ORDER BY transferred_at DESC LIMIT 500", *t_args
        )
    combined = [dict(r) for r in mov_rows] + [dict(r) for r in xfer_rows]
    combined.sort(key=lambda x: x.get("timestamp") or "", reverse=True)
    return combined[:1000]


@api.get("/reports/inventory-transactions")
async def report_inventory_transactions(
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    sku_id: Optional[str] = None,
    user: dict = Depends(get_user),
):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        conds: list = ["1=1"]
        args: list = []
        if date_from:
            args.append(date_from)
            conds.append(f"m.timestamp>=${len(args)}")
        if date_to:
            args.append(date_to + "T23:59:59")
            conds.append(f"m.timestamp<=${len(args)}")
        if sku_id:
            args.append(sku_id)
            conds.append(f"m.sku_id=${len(args)}")
        where = " AND ".join(conds)
        rows = await conn.fetch(
            f"SELECT m.id, m.type, m.timestamp, sk.sku_code, sk.name AS sku_name, "
            f"l.code AS location, l.zone, m.qty, m.ref, m.batch_no, m.expiry_date "
            f"FROM movements m JOIN skus sk ON sk.id=m.sku_id JOIN locations l ON l.id=m.location_id "
            f"WHERE {where} ORDER BY m.timestamp DESC LIMIT 1000", *args
        )
    return [dict(r) for r in rows]


@api.get("/reports/location-transfers")
async def report_location_transfers(
    date_from: Optional[str] = None,
    date_to: Optional[str] = None,
    user: dict = Depends(get_user),
):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        conds: list = ["1=1"]
        args: list = []
        if date_from:
            args.append(date_from)
            conds.append(f"transferred_at>=${len(args)}")
        if date_to:
            args.append(date_to + "T23:59:59")
            conds.append(f"transferred_at<=${len(args)}")
        where = " AND ".join(conds)
        rows = await conn.fetch(
            f"SELECT id, sku_code, sku_name, pallet_code, qty, bag_color, "
            f"from_code, to_code, notes, transferred_by, transferred_at, status, confirmed_at "
            f"FROM transfers WHERE {where} ORDER BY transferred_at DESC LIMIT 500", *args
        )
    return [dict(r) for r in rows]


@api.get("/reports/bin-utilization")
async def report_bin_utilization(user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        summary = await conn.fetchrow(
            "SELECT COUNT(*) AS total, "
            "SUM(CASE WHEN occupied>0 THEN 1 ELSE 0 END) AS occupied, "
            "SUM(CASE WHEN occupied=0 THEN 1 ELSE 0 END) AS empty, "
            "ROUND(SUM(CASE WHEN occupied>0 THEN 1 ELSE 0 END)*100.0/NULLIF(COUNT(*),0),1) AS pct "
            "FROM locations"
        )
        zones = await conn.fetch(
            "SELECT zone, COALESCE(zone_name,'') AS zone_name, COUNT(*) AS total, "
            "SUM(CASE WHEN occupied>0 THEN 1 ELSE 0 END) AS occupied, "
            "SUM(CASE WHEN occupied=0 THEN 1 ELSE 0 END) AS empty, "
            "ROUND(SUM(CASE WHEN occupied>0 THEN 1 ELSE 0 END)*100.0/NULLIF(COUNT(*),0),1) AS pct "
            "FROM locations GROUP BY zone,zone_name ORDER BY zone"
        )
    return {"summary": dict(summary), "zones": [dict(z) for z in zones]}


@api.get("/reports/warehouse-capacity")
async def report_warehouse_capacity(user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        zones = await conn.fetch(
            "SELECT zone, COALESCE(zone_name,'') AS zone_name, COUNT(*) AS total_locations, "
            "SUM(capacity) AS total_positions, SUM(occupied) AS occupied_positions, "
            "SUM(capacity-occupied) AS available_positions, "
            "ROUND(SUM(occupied)*100.0/NULLIF(SUM(capacity),0),1) AS utilization_pct "
            "FROM locations GROUP BY zone,zone_name ORDER BY zone"
        )
        totals = await conn.fetchrow(
            "SELECT SUM(capacity) AS total_positions, SUM(occupied) AS occupied_positions, "
            "SUM(capacity-occupied) AS available_positions, "
            "ROUND(SUM(occupied)*100.0/NULLIF(SUM(capacity),0),1) AS utilization_pct FROM locations"
        )
    return {"zones": [dict(r) for r in zones], "totals": dict(totals)}


@api.get("/reports/location-occupancy")
async def report_location_occupancy(zone: Optional[str] = None, user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        q = ("SELECT l.zone, COALESCE(l.zone_name,'') AS zone_name, l.row_label, "
             "l.lane_number, l.code, l.level, l.capacity, l.occupied, "
             "CASE WHEN l.occupied>0 THEN 'occupied' ELSE 'empty' END AS status, "
             "sk.sku_code, sk.name AS sku_name, COALESCE(SUM(s.qty),0) AS qty, "
             "MAX(s.expiry_date) AS expiry_date "
             "FROM locations l "
             "LEFT JOIN stock s ON s.location_id=l.id AND s.qty>0 "
             "LEFT JOIN skus sk ON sk.id=s.sku_id ")
        args: list = []
        if zone:
            args.append(zone)
            q += f"WHERE l.zone=${len(args)} "
        q += ("GROUP BY l.id,l.zone,l.zone_name,l.row_label,l.lane_number,l.code,"
              "l.level,l.capacity,l.occupied,sk.sku_code,sk.name "
              "ORDER BY l.zone,l.row_label,l.lane_number,l.level")
        rows = await conn.fetch(q, *args)
    return [dict(r) for r in rows]


@api.get("/reports/empty-locations")
async def report_empty_locations(zone: Optional[str] = None, user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        q = ("SELECT l.zone, COALESCE(l.zone_name,'') AS zone_name, l.code, "
             "l.row_label, l.lane_number, l.level, l.capacity, "
             "COALESCE(l.rack_type,'') AS rack_type "
             "FROM locations l WHERE l.occupied=0 ")
        args: list = []
        if zone:
            args.append(zone)
            q += f"AND l.zone=${len(args)} "
        q += "ORDER BY l.zone,l.row_label,l.lane_number,l.level"
        rows = await conn.fetch(q, *args)
    return [dict(r) for r in rows]


# ════════════════════════════════════════════════════
# ADMIN — UNITS OF MEASUREMENT
# ════════════════════════════════════════════════════

@api.get("/admin/units")
async def list_units(user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        rows = await conn.fetch("SELECT * FROM units_of_measurement ORDER BY code")
    return [dict(r) for r in rows]


class UnitBody(BaseModel):
    code: str
    name: str


@api.post("/admin/units")
async def create_unit(body: UnitBody, user: dict = Depends(require_role("admin"))):
    code = body.code.strip().upper()
    name = body.name.strip()
    if not code or not name:
        raise HTTPException(400, "Code and name are required")
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        existing = await conn.fetchrow("SELECT id FROM units_of_measurement WHERE code=$1", code)
        if existing:
            raise HTTPException(409, f"Unit '{code}' already exists")
        uid = str(uuid.uuid4())
        await conn.execute(
            "INSERT INTO units_of_measurement (id, code, name, created_at) VALUES ($1,$2,$3,$4)",
            uid, code, name, now_iso()
        )
        row = await conn.fetchrow("SELECT * FROM units_of_measurement WHERE id=$1", uid)
    return dict(row)


@api.delete("/admin/units/{unit_id}")
async def delete_unit(unit_id: str, user: dict = Depends(require_role("admin"))):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        row = await conn.fetchrow("SELECT code FROM units_of_measurement WHERE id=$1", unit_id)
        if not row:
            raise HTTPException(404, "Unit not found")
        # Prevent deleting if SKUs use this unit
        in_use = await conn.fetchval(
            "SELECT COUNT(*) FROM skus WHERE unit=$1", row["code"]
        )
        if in_use:
            raise HTTPException(409, f"Cannot delete '{row['code']}' — {in_use} SKU(s) use this unit")
        await conn.execute("DELETE FROM units_of_measurement WHERE id=$1", unit_id)
    return {"ok": True}


# ════════════════════════════════════════════════════
# ADMIN — SKU CATEGORIES
# ════════════════════════════════════════════════════

@api.get("/admin/categories")
async def list_categories(user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        rows = await conn.fetch("SELECT * FROM sku_categories ORDER BY name")
    return [dict(r) for r in rows]


class CategoryBody(BaseModel):
    name: str


@api.post("/admin/categories")
async def create_category(body: CategoryBody, user: dict = Depends(require_role("admin"))):
    name = body.name.strip()
    if not name:
        raise HTTPException(400, "Category name is required")
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        existing = await conn.fetchrow("SELECT id FROM sku_categories WHERE name=$1", name)
        if existing:
            raise HTTPException(409, f"Category '{name}' already exists")
        cid = str(uuid.uuid4())
        await conn.execute(
            "INSERT INTO sku_categories (id, name, created_at) VALUES ($1,$2,$3)",
            cid, name, now_iso()
        )
        row = await conn.fetchrow("SELECT * FROM sku_categories WHERE id=$1", cid)
    return dict(row)


@api.delete("/admin/categories/{cat_id}")
async def delete_category(cat_id: str, user: dict = Depends(require_role("admin"))):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        row = await conn.fetchrow("SELECT name FROM sku_categories WHERE id=$1", cat_id)
        if not row:
            raise HTTPException(404, "Category not found")
        in_use = await conn.fetchval(
            "SELECT COUNT(*) FROM skus WHERE category=$1", row["name"]
        )
        if in_use:
            raise HTTPException(409, f"Cannot delete '{row['name']}' — {in_use} SKU(s) use this category")
        await conn.execute("DELETE FROM sku_categories WHERE id=$1", cat_id)
    return {"ok": True}


# ════════════════════════════════════════════════════
# ADMIN — DATA MANAGEMENT
# ════════════════════════════════════════════════════

@api.post("/admin/purge-data")
async def purge_data(user: dict = Depends(require_role("admin"))):
    """Delete all transactional data (SKUs, stock, orders, movements).
    Zones, locations and users are preserved. Occupied counts are reset.
    Sets seed_locked flag so seed_data() does not re-populate on next restart."""
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        async with conn.transaction():
            await conn.execute("DELETE FROM shuttle_movements")
            await conn.execute("DELETE FROM movements")
            await conn.execute("DELETE FROM stock")
            await conn.execute("DELETE FROM inbound_items")
            await conn.execute("DELETE FROM inbound")
            await conn.execute("DELETE FROM outbound_items")
            await conn.execute("DELETE FROM outbound")
            await conn.execute("DELETE FROM skus")
            await conn.execute("UPDATE locations SET occupied = 0")
            await conn.execute(
                "INSERT INTO app_meta (key, value) VALUES ('seed_locked', '1') "
                "ON CONFLICT (key) DO UPDATE SET value = '1'"
            )
    return {"ok": True, "message": "All transactional data purged. Zones, locations and users preserved."}


# ════════════════════════════════════════════════════
# ADMIN — USER MANAGEMENT
# ════════════════════════════════════════════════════

@api.get("/admin/users")
async def list_users(user: dict = Depends(require_role("admin"))):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        rows = await conn.fetch(
            "SELECT id, email, name, role, created_at FROM users ORDER BY created_at"
        )
    return [dict(r) for r in rows]


@api.post("/admin/users")
async def create_user(body: RegisterIn, user: dict = Depends(require_role("admin"))):
    email = body.email.lower()
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        existing = await conn.fetchrow("SELECT id FROM users WHERE email = $1", email)
        if existing:
            raise HTTPException(400, "Email already registered")
        uid = str(uuid.uuid4())
        created = now_iso()
        await conn.execute(
            "INSERT INTO users (id, email, name, role, password_hash, created_at) VALUES ($1,$2,$3,$4,$5,$6)",
            uid, email, body.name, body.role, hash_pw(body.password), created,
        )
    return {"id": uid, "email": email, "name": body.name, "role": body.role, "created_at": created}


@api.delete("/admin/users/{user_id}")
async def delete_user(user_id: str, user: dict = Depends(require_role("admin"))):
    if user_id == user["id"]:
        raise HTTPException(400, "You cannot delete your own account")
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        target = await conn.fetchrow("SELECT id FROM users WHERE id = $1", user_id)
        if not target:
            raise HTTPException(404, "User not found")
        await conn.execute("DELETE FROM users WHERE id = $1", user_id)
    return {"ok": True, "message": "User removed"}


# ════════════════════════════════════════════════════
# SEED DATA
# ════════════════════════════════════════════════════

RACK_TYPES = {
    "A": {"depth": 4, "levels": 4, "weight_kg": 8000, "lanes": 13},
    "B": {"depth": 3, "levels": 4, "weight_kg": 12000, "lanes": 11},
    "C": {"depth": 5, "levels": 4, "weight_kg": 20000, "lanes": 15},
}

PLACEHOLDER_ZONES = [
    ("COLD-2", "Cold Storage 2", -18.0, True),
    ("COLD-3", "Cold Storage 3", -22.0, True),
    ("AMBIENT", "Ambient Warehouse", 22.0, True),
]


async def seed_data():
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        # ── Users ────────────────────────────────────────────────────
        users_seed = [
            ("ADMIN_EMAIL",   "ADMIN_PASSWORD",   "Admin User",         "admin"),
            ("MANAGER_EMAIL", "MANAGER_PASSWORD", "Warehouse Manager",  "manager"),
            ("OPERATOR_EMAIL","OPERATOR_PASSWORD","Floor Operator",      "operator"),
        ]
        for i in range(2, 10):
            em = os.environ.get(f"ADMIN{i}_EMAIL", "").lower()
            if not em:
                break
            users_seed.append((f"ADMIN{i}_EMAIL", f"ADMIN{i}_PASSWORD",
                                os.environ.get(f"ADMIN{i}_NAME", f"User {i}"),
                                os.environ.get(f"ADMIN{i}_ROLE", "operator")))
        for em_key, pw_key, name, role in users_seed:
            em = os.environ.get(em_key, "").lower()
            pw = os.environ.get(pw_key, "")
            if not em:
                continue
            existing = await conn.fetchrow("SELECT id, password_hash FROM users WHERE email=$1", em)
            if not existing:
                await conn.execute(
                    "INSERT INTO users (id, email, name, role, password_hash, created_at) VALUES ($1,$2,$3,$4,$5,$6)",
                    str(uuid.uuid4()), em, name, role, hash_pw(pw), now_iso(),
                )
            elif not verify_pw(pw, existing["password_hash"]):
                await conn.execute(
                    "UPDATE users SET password_hash=$1 WHERE id=$2", hash_pw(pw), existing["id"]
                )

        # ── Schema version ───────────────────────────────────────────
        meta = await conn.fetchrow("SELECT value FROM app_meta WHERE key='schema_version'")
        current_version = int((meta or {}).get("value", 0))
        if current_version < SCHEMA_VERSION:
            await conn.execute("DELETE FROM stock")
            await conn.execute("DELETE FROM movements")
            await conn.execute("DELETE FROM shuttle_movements")
            await conn.execute("DELETE FROM inbound_items")
            await conn.execute("DELETE FROM inbound")
            await conn.execute("DELETE FROM outbound_items")
            await conn.execute("DELETE FROM outbound")
            await conn.execute("DELETE FROM locations")
            await conn.execute("DELETE FROM skus")
            await conn.execute("DELETE FROM zones_meta")
            await conn.execute(
                "INSERT INTO app_meta (key, value) VALUES ('schema_version', $1) "
                "ON CONFLICT (key) DO UPDATE SET value=$1",
                str(SCHEMA_VERSION),
            )

        # ── Placeholder zones ────────────────────────────────────────
        cnt = await conn.fetchval("SELECT COUNT(*) FROM zones_meta")
        if cnt == 0:
            await conn.executemany(
                "INSERT INTO zones_meta (zone, name, temperature, placeholder) VALUES ($1,$2,$3,$4)",
                PLACEHOLDER_ZONES,
            )

        # ── COLD-1 locations ─────────────────────────────────────────
        cnt = await conn.fetchval("SELECT COUNT(*) FROM locations")
        if cnt == 0:
            locs = []
            plan = [
                ("A", "A", list(range(26, 39))),
                ("B", "B", list(range(39, 50))),
                ("C", "C", list(range(50, 65))),
            ]
            for row_label, rack_type, lane_numbers in plan:
                cfg = RACK_TYPES[rack_type]
                depth = cfg["depth"]
                levels = cfg["levels"]
                weight_kg = cfg["weight_kg"]
                for lane_num in lane_numbers:
                    for level in range(1, levels + 1):
                        for pos in range(1, depth + 1):
                            locs.append((
                                str(uuid.uuid4()),
                                f"COLD1-{row_label}-L{lane_num:02d}-LV{level}-P{pos:02d}",
                                "COLD-1", "Cold Storage 1", -20.0,
                                rack_type, row_label, lane_num, level, pos, depth, levels, weight_kg, 1, 0,
                            ))
            await conn.executemany(
                "INSERT INTO locations (id, code, zone, zone_name, temperature, rack_type, row_label, "
                "lane_number, level, position, depth, levels, weight_capacity_kg, capacity, occupied) "
                "VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10,$11,$12,$13,$14,$15)",
                locs,
            )

        # ── Shuttle zone locations ───────────────────────────────────
        shuttle_cnt = await conn.fetchval("SELECT COUNT(*) FROM locations WHERE zone=$1", SHUTTLE_ZONE)
        if shuttle_cnt == 0:
            shuttle_locs = []
            for lane in range(1, 8):
                for level in range(0, 5):
                    for depth in range(1, 51):
                        shuttle_locs.append((
                            str(uuid.uuid4()),
                            _shuttle_code(lane, level, depth),
                            SHUTTLE_ZONE, "Shuttle FIFO — Cold Storage", -20.0,
                            "shuttle", None, lane, level, depth, 50, 5, None, 1, 0,
                        ))
            # Insert in batches of 500 to avoid parameter limits
            for i in range(0, len(shuttle_locs), 500):
                await conn.executemany(
                    "INSERT INTO locations (id, code, zone, zone_name, temperature, rack_type, row_label, "
                    "lane_number, level, position, depth, levels, weight_capacity_kg, capacity, occupied) "
                    "VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10,$11,$12,$13,$14,$15)",
                    shuttle_locs[i:i+500],
                )

        # ── SKUs + initial stock ─────────────────────────────────────
        seed_locked = await conn.fetchval("SELECT value FROM app_meta WHERE key='seed_locked'")
        sku_cnt = await conn.fetchval("SELECT COUNT(*) FROM skus")
        if sku_cnt == 0 and not seed_locked:
            catalog = [
                ("FRZ-MEAT-001","Frozen Beef Sirloin 20kg","Frozen Meat","PLT",480.00,4),
                ("FRZ-MEAT-002","Frozen Chicken Breast 25kg","Frozen Meat","PLT",320.00,5),
                ("FRZ-MEAT-003","Frozen Pork Loin 22kg","Frozen Meat","PLT",380.00,4),
                ("FRZ-MEAT-004","Frozen Lamb Chops 18kg","Frozen Meat","PLT",540.00,3),
                ("FRZ-SEAF-001","Frozen Atlantic Salmon 20kg","Seafood","PLT",720.00,3),
                ("FRZ-SEAF-002","Frozen Tiger Prawns 15kg","Seafood","PLT",850.00,2),
                ("FRZ-SEAF-003","Frozen Tuna Steaks 18kg","Seafood","PLT",920.00,3),
                ("FRZ-DAIRY-001","Frozen Butter Blocks 25kg","Dairy","PLT",220.00,6),
                ("FRZ-DAIRY-002","Vanilla Ice Cream Tubs 20L","Ice Cream","PLT",180.00,8),
                ("FRZ-DAIRY-003","Chocolate Ice Cream Tubs 20L","Ice Cream","PLT",195.00,6),
                ("FRZ-DAIRY-004","Strawberry Ice Cream 20L","Ice Cream","PLT",200.00,5),
                ("FRZ-VEG-001","Frozen Mixed Vegetables 15kg","Vegetables","PLT",120.00,8),
                ("FRZ-VEG-002","Frozen Sweet Corn 15kg","Vegetables","PLT",95.00,10),
                ("FRZ-VEG-003","Frozen Spinach Blocks 12kg","Vegetables","PLT",110.00,6),
                ("FRZ-VEG-004","Frozen Green Peas 15kg","Vegetables","PLT",105.00,8),
                ("FRZ-FRUIT-001","Frozen Strawberries 12kg","Fruits","PLT",240.00,5),
                ("FRZ-FRUIT-002","Frozen Blueberries 10kg","Fruits","PLT",290.00,4),
                ("FRZ-FRUIT-003","Frozen Mango Chunks 12kg","Fruits","PLT",220.00,5),
                ("FRZ-DOUGH-001","Frozen Pizza Dough Balls","Bakery","PLT",140.00,6),
                ("FRZ-DOUGH-002","Frozen Croissant Dough","Bakery","PLT",165.00,4),
                ("FRZ-READY-001","Frozen Lasagna Trays","Ready Meals","PLT",280.00,5),
                ("FRZ-READY-002","Frozen Chicken Curry Trays","Ready Meals","PLT",260.00,4),
                ("FRZ-PROC-001","Frozen Chicken Nuggets 12kg","Processed","PLT",175.00,8),
                ("FRZ-PROC-002","Frozen French Fries 15kg","Processed","PLT",130.00,12),
                ("FRZ-PROC-003","Frozen Spring Rolls 10kg","Processed","PLT",155.00,6),
            ]
            sku_rows_data = []
            ts = now_iso()
            for code, name, cat, unit, price, reorder in catalog:
                color, wpb, bpp, dims = _packaging_for(cat)
                sku_rows_data.append((
                    str(uuid.uuid4()), code, name, cat, unit, price, reorder, 0, ts,
                    color, _to_decimal(wpb), bpp, dims,
                ))
            await conn.executemany(
                "INSERT INTO skus (id, sku_code, name, category, unit, unit_price, reorder_level, total_stock, created_at, "
                "bag_color, weight_per_bag, bags_per_pallet, dimensions) "
                "VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10,$11,$12,$13)",
                sku_rows_data,
            )
            # Fetch inserted skus and COLD-1 locations for initial stock distribution
            skus_db = await conn.fetch("SELECT * FROM skus")
            all_locs = list(await conn.fetch("SELECT * FROM locations WHERE zone='COLD-1'"))
            random.shuffle(all_locs)
            target_fill = int(len(all_locs) * 0.58)
            placed_total = 0
            for sku in skus_db:
                pallets = random.randint(8, 28)
                placed = 0
                while placed < pallets and all_locs:
                    loc = all_locs.pop()
                    if loc["occupied"] >= loc["capacity"]:
                        continue
                    ts2 = now_iso()
                    mfg = (datetime.now(timezone.utc) - timedelta(days=random.randint(30,180))).date().isoformat()
                    exp = (datetime.now(timezone.utc) + timedelta(days=random.randint(15,365))).date().isoformat()
                    recv = (datetime.now(timezone.utc) - timedelta(days=random.randint(1,13))).isoformat()
                    await conn.execute(
                        "INSERT INTO stock (id, sku_id, location_id, qty, batch_no, manufacture_date, expiry_date, received_date, pallet_status, original_qty) "
                        "VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10)",
                        str(uuid.uuid4()), sku["id"], loc["id"], 1,
                        f"B{random.randint(1000,9999)}-{sku['sku_code'][-3:]}",
                        mfg, exp, recv, "full", 1,
                    )
                    await conn.execute("UPDATE locations SET occupied=1 WHERE id=$1", loc["id"])
                    mv_ts = (datetime.now(timezone.utc) - timedelta(days=random.randint(1,13), hours=random.randint(0,23))).isoformat()
                    await conn.execute(
                        "INSERT INTO movements (id, type, sku_id, location_id, qty, ref, timestamp) VALUES ($1,$2,$3,$4,$5,$6,$7)",
                        str(uuid.uuid4()), "in", sku["id"], loc["id"], 1,
                        f"PO-INIT-{random.randint(1000,9999)}", mv_ts,
                    )
                    placed += 1
                    placed_total += 1
                    if placed_total >= target_fill:
                        break
                await conn.execute("UPDATE skus SET total_stock=total_stock+$1 WHERE id=$2", placed, sku["id"])
                if placed_total >= target_fill:
                    break

            # Historical outbound
            skus_final = await conn.fetch("SELECT id FROM skus")
            for _ in range(35):
                sku = random.choice(skus_final)
                s = await conn.fetchrow("SELECT * FROM stock WHERE sku_id=$1 AND qty > 0 LIMIT 1", sku["id"])
                if not s:
                    continue
                mv_ts = (datetime.now(timezone.utc) - timedelta(days=random.randint(0,13), hours=random.randint(0,23))).isoformat()
                await conn.execute("UPDATE stock SET qty=qty-1 WHERE id=$1", s["id"])
                await conn.execute("UPDATE locations SET occupied=GREATEST(0,occupied-1) WHERE id=$1", s["location_id"])
                await conn.execute("UPDATE skus SET total_stock=GREATEST(0,total_stock-1) WHERE id=$1", sku["id"])
                await conn.execute(
                    "INSERT INTO movements (id, type, sku_id, location_id, qty, ref, timestamp) VALUES ($1,$2,$3,$4,$5,$6,$7)",
                    str(uuid.uuid4()), "out", sku["id"], s["location_id"], 1,
                    f"SO-INIT-{random.randint(1000,9999)}", mv_ts,
                )

        # ── Backfill packaging details for SKUs missing them ─────────
        # Columns (bag_color/weight_per_bag/bags_per_pallet/dimensions) were
        # added after some SKUs were created, leaving them NULL. Fill any gaps
        # with sensible category-based defaults so the catalog stays complete.
        # Bag Color is user-optional, so it is deliberately NOT backfilled here —
        # only the still-required packaging fields are filled for legacy SKUs.
        missing = await conn.fetch(
            "SELECT id, category FROM skus WHERE weight_per_bag IS NULL "
            "OR bags_per_pallet IS NULL OR dimensions IS NULL"
        )
        for s in missing:
            _, wpb, bpp, dims = _packaging_for(s["category"])
            await conn.execute(
                "UPDATE skus SET weight_per_bag=COALESCE(weight_per_bag,$1), "
                "bags_per_pallet=COALESCE(bags_per_pallet,$2), "
                "dimensions=COALESCE(dimensions,$3) WHERE id=$4",
                _to_decimal(wpb), bpp, dims, s["id"],
            )

        # ── Backfill stock.bag_color for rows received before the column ──
        # existed. Prefer the actual colour recorded on the originating inbound
        # line (matched via PO ref + sku + location). To stay deterministic we
        # only backfill when the candidate lines agree on a SINGLE colour;
        # ambiguous matches (e.g. a PO with mixed colours for the same
        # SKU/location, or a reused PO number) are left NULL rather than guessed.
        await conn.execute(
            "UPDATE stock s SET bag_color = c.color FROM ("
            "  SELECT s2.id AS stock_id, MIN(ii.bag_color) AS color "
            "  FROM stock s2 "
            "  JOIN inbound i ON s2.ref = i.po_number "
            "  JOIN inbound_items ii ON ii.inbound_id = i.id "
            "    AND ii.sku_id = s2.sku_id AND ii.location_id = s2.location_id "
            "  WHERE s2.bag_color IS NULL AND ii.bag_color IS NOT NULL "
            "  GROUP BY s2.id HAVING COUNT(DISTINCT ii.bag_color) = 1"
            ") c WHERE s.id = c.stock_id"
        )
        # Remaining gaps fall back to the SKU's default colour, if it has one.
        await conn.execute(
            "UPDATE stock s SET bag_color = sk.bag_color "
            "FROM skus sk WHERE s.bag_color IS NULL AND sk.id = s.sku_id "
            "AND sk.bag_color IS NOT NULL"
        )

        # ── Sample inbound orders ────────────────────────────────────
        inb_cnt = await conn.fetchval("SELECT COUNT(*) FROM inbound")
        if inb_cnt == 0 and not seed_locked:
            skus_db = await conn.fetch("SELECT id FROM skus")
            free_locs = list(await conn.fetch(
                "SELECT id FROM locations WHERE occupied < capacity ORDER BY id"
            ))
            suppliers = ["Arctic Foods Ltd","FrostChain Suppliers","PolarFresh Co","Glacier Distributors","IceVault Logistics"]
            for i in range(8):
                oid2 = str(uuid.uuid4())
                picks = random.sample(list(skus_db), k=min(random.randint(2,4), len(skus_db)))
                crat = (datetime.now(timezone.utc) - timedelta(days=random.randint(0,5))).isoformat()
                exp_d = (datetime.now(timezone.utc) + timedelta(days=random.randint(1,10))).date().isoformat()
                status = random.choice(["pending","pending","completed"])
                await conn.execute(
                    "INSERT INTO inbound (id, po_number, supplier, expected_date, status, created_at, created_by) "
                    "VALUES ($1,$2,$3,$4,$5,$6,$7)",
                    oid2, f"PO-2026-{1000+i:04d}", random.choice(suppliers), exp_d, status, crat, "Admin User",
                )
                for s in picks:
                    for _ in range(random.randint(1,3)):
                        if not free_locs:
                            break
                        loc = free_locs.pop()
                        await conn.execute(
                            "INSERT INTO inbound_items (id, inbound_id, sku_id, qty, location_id, barcode, putaway_confirmed) "
                            "VALUES ($1,$2,$3,$4,$5,$6,$7)",
                            str(uuid.uuid4()), oid2, s["id"], 1, loc["id"],
                            f"PLT-PO-2026-{1000+i:04d}-{random.randint(1,99):02d}", False,
                        )

        # ── Sample outbound orders ───────────────────────────────────
        out_cnt = await conn.fetchval("SELECT COUNT(*) FROM outbound")
        if out_cnt == 0 and not seed_locked:
            skus_db = await conn.fetch("SELECT id FROM skus")
            customers = ["Metro Supermarkets","FreshMart Chain","ColdLink Retail","OmegaFoods","GroceryPro Inc"]
            statuses = ["pending","picking","packing","shipped"]
            for i in range(10):
                oid3 = str(uuid.uuid4())
                picks = random.sample(list(skus_db), k=min(random.randint(2,4), len(skus_db)))
                crat = (datetime.now(timezone.utc) - timedelta(days=random.randint(0,5))).isoformat()
                await conn.execute(
                    "INSERT INTO outbound (id, so_number, customer, status, created_at, created_by) VALUES ($1,$2,$3,$4,$5,$6)",
                    oid3, f"SO-2026-{2000+i:04d}", random.choice(customers),
                    random.choice(statuses), crat, "Admin User",
                )
                for s in picks:
                    await conn.execute(
                        "INSERT INTO outbound_items (id, outbound_id, sku_id, qty) VALUES ($1,$2,$3,$4)",
                        str(uuid.uuid4()), oid3, s["id"], random.randint(1,4),
                    )


# ---------- Startup / Shutdown ----------
@app.on_event("startup")
async def on_startup():
    await _db.init_db()
    await seed_data()


@app.on_event("shutdown")
async def on_shutdown():
    await _db.close_pool()


# ════════════════════════════════════════════════════
# CUSTOMERS & VENDORS
# ════════════════════════════════════════════════════

class CustomerIn(BaseModel):
    code: str
    name: str
    contact_person: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None
    address: Optional[str] = None
    city: Optional[str] = None
    country: Optional[str] = None
    notes: Optional[str] = None

class VendorIn(BaseModel):
    code: str
    name: str
    contact_person: Optional[str] = None
    email: Optional[str] = None
    phone: Optional[str] = None
    address: Optional[str] = None
    city: Optional[str] = None
    country: Optional[str] = None
    notes: Optional[str] = None

class OrderItemIn(BaseModel):
    sku_id: str
    qty: int
    unit_cost: float

class SaleItemIn(BaseModel):
    sku_id: str
    qty: int
    unit_price: float

class PurchaseIn(BaseModel):
    vendor_id: str
    order_date: str
    expected_date: Optional[str] = None
    notes: Optional[str] = None
    items: List[OrderItemIn]

class SaleIn(BaseModel):
    customer_id: str
    order_date: str
    expected_date: Optional[str] = None
    notes: Optional[str] = None
    items: List[SaleItemIn]


# ── Customers ───────────────────────────────────────

@api.get("/customers")
async def list_customers(user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        rows = await conn.fetch(
            "SELECT c.*, "
            "(SELECT COUNT(*) FROM sales_orders WHERE customer_id=c.id) AS order_count, "
            "(SELECT COALESCE(SUM(total_amount),0) FROM sales_orders WHERE customer_id=c.id) AS total_sales "
            "FROM customers c ORDER BY c.name"
        )
    return _db.rl(rows)


@api.post("/customers")
async def create_customer(body: CustomerIn, user: dict = Depends(require_role("manager", "admin"))):
    pool = await _db.get_pool()
    cid = str(uuid.uuid4())
    now = datetime.utcnow().isoformat()
    async with pool.acquire() as conn:
        existing = await conn.fetchrow("SELECT id FROM customers WHERE code=$1", body.code.upper())
        if existing:
            raise HTTPException(400, f"Customer code {body.code.upper()} already exists")
        row = await conn.fetchrow(
            "INSERT INTO customers (id,code,name,contact_person,email,phone,address,city,country,notes,created_at) "
            "VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10,$11) RETURNING *",
            cid, body.code.upper(), body.name, body.contact_person, body.email,
            body.phone, body.address, body.city, body.country, body.notes, now,
        )
    return _db.r(row)


@api.put("/customers/{cid}")
async def update_customer(cid: str, body: CustomerIn, user: dict = Depends(require_role("manager", "admin"))):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        row = await conn.fetchrow(
            "UPDATE customers SET code=$1,name=$2,contact_person=$3,email=$4,phone=$5,"
            "address=$6,city=$7,country=$8,notes=$9 WHERE id=$10 RETURNING *",
            body.code.upper(), body.name, body.contact_person, body.email,
            body.phone, body.address, body.city, body.country, body.notes, cid,
        )
    if not row:
        raise HTTPException(404, "Customer not found")
    return _db.r(row)


@api.delete("/customers/{cid}")
async def delete_customer(cid: str, user: dict = Depends(require_role("manager", "admin"))):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        orders = await conn.fetchval("SELECT COUNT(*) FROM sales_orders WHERE customer_id=$1", cid)
        if orders:
            raise HTTPException(409, f"Cannot delete: {orders} sales order(s) linked to this customer")
        res = await conn.execute("DELETE FROM customers WHERE id=$1", cid)
    if res == "DELETE 0":
        raise HTTPException(404, "Customer not found")
    return {"ok": True}


# ── Vendors ─────────────────────────────────────────

@api.get("/vendors")
async def list_vendors(user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        rows = await conn.fetch(
            "SELECT v.*, "
            "(SELECT COUNT(*) FROM purchase_orders WHERE vendor_id=v.id) AS order_count, "
            "(SELECT COALESCE(SUM(total_amount),0) FROM purchase_orders WHERE vendor_id=v.id) AS total_purchases "
            "FROM vendors v ORDER BY v.name"
        )
    return _db.rl(rows)


@api.post("/vendors")
async def create_vendor(body: VendorIn, user: dict = Depends(require_role("manager", "admin"))):
    pool = await _db.get_pool()
    vid = str(uuid.uuid4())
    now = datetime.utcnow().isoformat()
    async with pool.acquire() as conn:
        existing = await conn.fetchrow("SELECT id FROM vendors WHERE code=$1", body.code.upper())
        if existing:
            raise HTTPException(400, f"Vendor code {body.code.upper()} already exists")
        row = await conn.fetchrow(
            "INSERT INTO vendors (id,code,name,contact_person,email,phone,address,city,country,notes,created_at) "
            "VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10,$11) RETURNING *",
            vid, body.code.upper(), body.name, body.contact_person, body.email,
            body.phone, body.address, body.city, body.country, body.notes, now,
        )
    return _db.r(row)


@api.put("/vendors/{vid}")
async def update_vendor(vid: str, body: VendorIn, user: dict = Depends(require_role("manager", "admin"))):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        row = await conn.fetchrow(
            "UPDATE vendors SET code=$1,name=$2,contact_person=$3,email=$4,phone=$5,"
            "address=$6,city=$7,country=$8,notes=$9 WHERE id=$10 RETURNING *",
            body.code.upper(), body.name, body.contact_person, body.email,
            body.phone, body.address, body.city, body.country, body.notes, vid,
        )
    if not row:
        raise HTTPException(404, "Vendor not found")
    return _db.r(row)


@api.delete("/vendors/{vid}")
async def delete_vendor(vid: str, user: dict = Depends(require_role("manager", "admin"))):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        orders = await conn.fetchval("SELECT COUNT(*) FROM purchase_orders WHERE vendor_id=$1", vid)
        if orders:
            raise HTTPException(409, f"Cannot delete: {orders} purchase order(s) linked to this vendor")
        res = await conn.execute("DELETE FROM vendors WHERE id=$1", vid)
    if res == "DELETE 0":
        raise HTTPException(404, "Vendor not found")
    return {"ok": True}


# ── Purchase Orders ──────────────────────────────────

def _po_number():
    from random import randint
    return f"PO-{datetime.utcnow().strftime('%Y%m')}-{randint(1000,9999)}"


@api.get("/purchases")
async def list_purchases(user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        rows = await conn.fetch(
            "SELECT p.*, v.name AS vendor_name, v.code AS vendor_code, "
            "(SELECT COUNT(*) FROM purchase_order_items WHERE po_id=p.id) AS item_count "
            "FROM purchase_orders p JOIN vendors v ON v.id=p.vendor_id "
            "ORDER BY p.created_at DESC"
        )
    return _db.rl(rows)


@api.get("/purchases/{po_id}")
async def get_purchase(po_id: str, user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        po = await conn.fetchrow(
            "SELECT p.*, v.name AS vendor_name, v.code AS vendor_code "
            "FROM purchase_orders p JOIN vendors v ON v.id=p.vendor_id WHERE p.id=$1", po_id
        )
        if not po:
            raise HTTPException(404, "Purchase order not found")
        items = await conn.fetch(
            "SELECT * FROM purchase_order_items WHERE po_id=$1 ORDER BY id", po_id
        )
    result = _db.r(po)
    result["items"] = _db.rl(items)
    return result


@api.post("/purchases")
async def create_purchase(body: PurchaseIn, user: dict = Depends(require_role("manager", "admin"))):
    pool = await _db.get_pool()
    po_id = str(uuid.uuid4())
    po_number = _po_number()
    now = datetime.utcnow().isoformat()
    async with pool.acquire() as conn:
        vendor = await conn.fetchrow("SELECT id FROM vendors WHERE id=$1", body.vendor_id)
        if not vendor:
            raise HTTPException(404, "Vendor not found")
        skus = {r["id"]: r for r in await conn.fetch("SELECT id,sku_code,name FROM skus WHERE id=ANY($1)", [i.sku_id for i in body.items])}
        total = sum(i.qty * i.unit_cost for i in body.items)
        await conn.execute(
            "INSERT INTO purchase_orders (id,po_number,vendor_id,status,order_date,expected_date,total_amount,notes,created_by,created_at) "
            "VALUES ($1,$2,$3,'draft',$4,$5,$6,$7,$8,$9)",
            po_id, po_number, body.vendor_id, body.order_date, body.expected_date,
            total, body.notes, user["email"], now,
        )
        for item in body.items:
            sku = skus.get(item.sku_id)
            if not sku:
                raise HTTPException(404, f"SKU {item.sku_id} not found")
            await conn.execute(
                "INSERT INTO purchase_order_items (id,po_id,sku_id,sku_code,sku_name,qty,unit_cost,total_cost) "
                "VALUES ($1,$2,$3,$4,$5,$6,$7,$8)",
                str(uuid.uuid4()), po_id, item.sku_id, sku["sku_code"], sku["name"],
                item.qty, item.unit_cost, item.qty * item.unit_cost,
            )
    return {"id": po_id, "po_number": po_number}


@api.put("/purchases/{po_id}/status")
async def update_purchase_status(po_id: str, body: dict, user: dict = Depends(require_role("manager", "admin"))):
    status = body.get("status")
    if status not in ("draft", "confirmed", "received", "cancelled"):
        raise HTTPException(400, "Invalid status")
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        res = await conn.execute("UPDATE purchase_orders SET status=$1 WHERE id=$2", status, po_id)
    if res == "UPDATE 0":
        raise HTTPException(404, "Purchase order not found")
    return {"ok": True}


@api.delete("/purchases/{po_id}")
async def delete_purchase(po_id: str, user: dict = Depends(require_role("manager", "admin"))):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        po = await conn.fetchrow("SELECT status FROM purchase_orders WHERE id=$1", po_id)
        if not po:
            raise HTTPException(404, "Purchase order not found")
        if po["status"] not in ("draft", "cancelled"):
            raise HTTPException(409, "Only draft or cancelled orders can be deleted")
        await conn.execute("DELETE FROM purchase_orders WHERE id=$1", po_id)
    return {"ok": True}


# ── Sales Orders ─────────────────────────────────────

def _so_number():
    from random import randint
    return f"SO-{datetime.utcnow().strftime('%Y%m')}-{randint(1000,9999)}"


@api.get("/sales")
async def list_sales(user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        rows = await conn.fetch(
            "SELECT s.*, c.name AS customer_name, c.code AS customer_code, "
            "(SELECT COUNT(*) FROM sales_order_items WHERE so_id=s.id) AS item_count "
            "FROM sales_orders s JOIN customers c ON c.id=s.customer_id "
            "ORDER BY s.created_at DESC"
        )
    return _db.rl(rows)


@api.get("/sales/{so_id}")
async def get_sale(so_id: str, user: dict = Depends(get_user)):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        so = await conn.fetchrow(
            "SELECT s.*, c.name AS customer_name, c.code AS customer_code "
            "FROM sales_orders s JOIN customers c ON c.id=s.customer_id WHERE s.id=$1", so_id
        )
        if not so:
            raise HTTPException(404, "Sales order not found")
        items = await conn.fetch(
            "SELECT * FROM sales_order_items WHERE so_id=$1 ORDER BY id", so_id
        )
    result = _db.r(so)
    result["items"] = _db.rl(items)
    return result


@api.post("/sales")
async def create_sale(body: SaleIn, user: dict = Depends(require_role("manager", "admin"))):
    pool = await _db.get_pool()
    so_id = str(uuid.uuid4())
    so_number = _so_number()
    now = datetime.utcnow().isoformat()
    async with pool.acquire() as conn:
        customer = await conn.fetchrow("SELECT id FROM customers WHERE id=$1", body.customer_id)
        if not customer:
            raise HTTPException(404, "Customer not found")
        skus = {r["id"]: r for r in await conn.fetch("SELECT id,sku_code,name FROM skus WHERE id=ANY($1)", [i.sku_id for i in body.items])}
        total = sum(i.qty * i.unit_price for i in body.items)
        await conn.execute(
            "INSERT INTO sales_orders (id,so_number,customer_id,status,order_date,expected_date,total_amount,notes,created_by,created_at) "
            "VALUES ($1,$2,$3,'draft',$4,$5,$6,$7,$8,$9)",
            so_id, so_number, body.customer_id, body.order_date, body.expected_date,
            total, body.notes, user["email"], now,
        )
        for item in body.items:
            sku = skus.get(item.sku_id)
            if not sku:
                raise HTTPException(404, f"SKU {item.sku_id} not found")
            await conn.execute(
                "INSERT INTO sales_order_items (id,so_id,sku_id,sku_code,sku_name,qty,unit_price,total_price) "
                "VALUES ($1,$2,$3,$4,$5,$6,$7,$8)",
                str(uuid.uuid4()), so_id, item.sku_id, sku["sku_code"], sku["name"],
                item.qty, item.unit_price, item.qty * item.unit_price,
            )
    return {"id": so_id, "so_number": so_number}


@api.put("/sales/{so_id}/status")
async def update_sale_status(so_id: str, body: dict, user: dict = Depends(require_role("manager", "admin"))):
    status = body.get("status")
    if status not in ("draft", "confirmed", "shipped", "cancelled"):
        raise HTTPException(400, "Invalid status")
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        res = await conn.execute("UPDATE sales_orders SET status=$1 WHERE id=$2", status, so_id)
    if res == "UPDATE 0":
        raise HTTPException(404, "Sales order not found")
    return {"ok": True}


@api.delete("/sales/{so_id}")
async def delete_sale(so_id: str, user: dict = Depends(require_role("manager", "admin"))):
    pool = await _db.get_pool()
    async with pool.acquire() as conn:
        so = await conn.fetchrow("SELECT status FROM sales_orders WHERE id=$1", so_id)
        if not so:
            raise HTTPException(404, "Sales order not found")
        if so["status"] not in ("draft", "cancelled"):
            raise HTTPException(409, "Only draft or cancelled orders can be deleted")
        await conn.execute("DELETE FROM sales_orders WHERE id=$1", so_id)
    return {"ok": True}


# ---------- Mount ----------
app.include_router(api)

origins_env = os.environ.get("CORS_ORIGINS", "*")
allowed = [o.strip() for o in origins_env.split(",")] if origins_env != "*" else ["*"]
app.add_middleware(
    CORSMiddleware,
    allow_origins=allowed,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")


@api.get("/")
async def root():
    return {"app": "FrostCore API", "status": "ok", "db": "postgresql"}
